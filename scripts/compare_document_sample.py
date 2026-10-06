"""Preserve exact semantic content and record every question's before/after layout."""
import argparse
import hashlib
import json
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn

from nas_filetools.document_layout_checks import compare_docx_content, paragraph_locations
from nas_filetools.document_sample import item_groups


def boundaries(path, items, has_footer):
    with fitz.open(path) as pdf:
        rows = paragraph_locations(pdf, items, expect_document_footer=has_footer)
    questions = []
    for start, end, group in item_groups(items):
        if group[0][0] != "Question":
            continue
        selected = rows[start:end]
        first = selected[0]
        questions.append({"number": group[0][1].split(".", 1)[0],
                          "pages": sorted({page for row in selected for page in row["pages"]}),
                          "stem_opening_lines": sum(line["page"] == first["pages"][0] for line in first["lines"]),
                          "last_option_pages": selected[-2]["pages"], "source_pages": selected[-1]["pages"],
                          "source_follows": selected[-2]["pages"][-1] == selected[-1]["pages"][0]})
    return questions


def compare(before, delivery, rendered, output):
    before, delivery, rendered, output = map(Path, (before, delivery, rendered, output))
    result = compare_docx_content(before / "delivery/sample.docx", delivery / "sample.docx")
    hashes = {}
    for name in ("structured.json", "issues.md", "original-ocr.md"):
        old, new = (before / "delivery" / name).read_bytes(), (delivery / name).read_bytes()
        if old != new:
            raise ValueError("UNCHANGED_AUDIT_OR_ORIGINAL_REQUIRED: " + name)
        hashes[name] = hashlib.sha256(new).hexdigest()
    result["unchanged_audit_and_original_hashes"] = hashes
    old_word = Document(before / "delivery/sample.docx")
    old_items = [(p.style.name, p.text) for p in old_word.paragraphs]
    old_has_footer = any("PAGE" in field.get(qn("w:instr"), "")
                         for section in old_word.sections
                         for field in section.footer._element.iter(qn("w:fldSimple")))
    new_items = [(p.style.name, p.text) for p in Document(delivery / "sample.docx").paragraphs]
    for label, old_path, new_path in [("word", before / "rendered/sample.pdf", rendered / "sample.pdf"),
                                     ("pdf", before / "delivery/sample.pdf", delivery / "sample.pdf")]:
        old = boundaries(old_path, old_items, old_has_footer)
        new = boundaries(new_path, new_items, True)
        if [q["number"] for q in old] != [q["number"] for q in new]:
            raise ValueError("QUESTION_ORDER_CHANGED")
        result[label] = {"before_source_orphans": sum(not q["source_follows"] for q in old),
                         "after_source_orphans": sum(not q["source_follows"] for q in new),
                         "questions": [{"ordinal": n, "before": a, "after": b} for n, (a, b) in enumerate(zip(old, new), 1)]}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(json.dumps({k: v for k, v in result.items() if k not in ("word", "pdf")}, ensure_ascii=False))
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    for name in ("before", "delivery", "rendered", "output"):
        parser.add_argument("--" + name, required=True, type=Path)
    args = parser.parse_args()
    compare(args.before, args.delivery, args.rendered, args.output)
