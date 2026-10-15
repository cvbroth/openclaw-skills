"""Build the versioned Pandoc reference, never export body content with python-docx."""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

p = Path(__file__).resolve().parents[1] / "src/nas_filetools/standalone/reference.docx"
doc = Document()
s = doc.sections[0]
s.page_width = Mm(210)
s.page_height = Mm(297)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Mm(18)
s.header_distance = s.footer_distance = Mm(8)
for name in [
    "Normal",
    "Body Text",
    "First Paragraph",
    "Compact",
    "Block Text",
    "Source Code",
    "Verbatim Char",
    "Hyperlink",
    "Table",
]:
    if name not in doc.styles:
        kind = (
            WD_STYLE_TYPE.TABLE
            if name == "Table"
            else WD_STYLE_TYPE.CHARACTER
            if name in {"Verbatim Char", "Hyperlink"}
            else WD_STYLE_TYPE.PARAGRAPH
        )
        doc.styles.add_style(name, kind)
for style in doc.styles:
    if style.type in {WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.TABLE}:
        style.font.name = "Droid Sans Fallback"
        style.font.size = Pt(10.5)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn("w:" + k), "Droid Sans Fallback")
    if style.type == WD_STYLE_TYPE.PARAGRAPH:
        f = style.paragraph_format
        f.line_spacing = 1.35
        f.space_before = Pt(0)
        f.space_after = Pt(5)
        f.widow_control = True
for level, size in enumerate([18, 16, 14, 12, 11, 10.5], 1):
    style = doc.styles[f"Heading {level}"]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True
for name in ["Header", "Footer"]:
    doc.styles[name].font.size = Pt(9)
doc.styles["Block Text"].paragraph_format.left_indent = Mm(6)
doc.styles["Source Code"].paragraph_format.space_after = Pt(2)
table = doc.styles["Table"]
table.font.size = Pt(9)
props = OxmlElement("w:tblPr")
borders = OxmlElement("w:tblBorders")
for side in ["top", "bottom", "left", "right", "insideH", "insideV"]:
    edge = OxmlElement("w:" + side)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "4")
    edge.set(qn("w:color"), "64748B")
    borders.append(edge)
props.append(borders)
table.element.append(props)
cell_props = OxmlElement("w:tcPr")
vertical = OxmlElement("w:vAlign")
vertical.set(qn("w:val"), "center")
cell_props.append(vertical)
table.element.append(cell_props)
# Align repeated document furniture independently of body.
s.header.paragraphs[0].alignment = 1
s.footer.paragraphs[0].alignment = 1
# Pandoc takes styles and section properties, not this sample body.
doc.add_paragraph("FileTools Pandoc reference v1.0.0")
doc.save(p)
print(p)
