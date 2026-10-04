"""Generic temporary Python jobs with verified output copies, not Excel business tools."""
import json
import os
import re
import shutil
import sys
import time
import uuid
from pathlib import Path

from .catalog import digest_file, identify
from .contracts import Fault, strict


def request(store, identity, params):
    strict(params, ("attachment_id", "code", "outputs", "checks", "description"),
           ("attachment_id", "code", "outputs", "description"))
    if (not isinstance(params["code"], str) or not 1 <= len(params["code"]) <= 64000 or
            not isinstance(params["description"], str) or not 1 <= len(params["description"]) <= 1000 or
            not isinstance(params["outputs"], list) or not 1 <= len(params["outputs"]) <= 8 or
            any(not isinstance(n, str) or Path(n).name != n or any(c in n for c in "\\/:\0") or
                n in ("operation.py", "stdout.log", "stderr.log", "sources.json", "content.md", "input.xlsx")
                for n in params["outputs"])):
        raise Fault("INVALID_SCRIPT")
    checks = params.get("checks", [])
    if not isinstance(checks, list) or len(checks) > 100:
        raise Fault("INVALID_SCRIPT")
    for check in checks:
        strict(check, ("file", "sheet", "cell", "equals", "min_rows", "min_columns"), ("file", "sheet"))
        if check["file"] not in params["outputs"] or not isinstance(check["sheet"], str):
            raise Fault("INVALID_SCRIPT")
        if not 1 <= len(check["sheet"]) <= 31 or any(type(check[k]) is not int or not 1 <= check[k] <= 1048576
            for k in ("min_rows", "min_columns") if k in check):
            raise Fault("INVALID_SCRIPT")
        if "cell" in check and (not isinstance(check["cell"], str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", check["cell"])):
            raise Fault("INVALID_SCRIPT")
        if "equals" in check and (type(check["equals"]) not in (str, int, float, bool, type(None)) or
                                   isinstance(check["equals"], str) and len(check["equals"]) > 8000):
            raise Fault("INVALID_SCRIPT")
    if store.limits.script_isolation == "landlock" and sys.platform != "linux":
        raise Fault("SCRIPT_ISOLATION_UNAVAILABLE")
    attachment = store.attachment(identity, params["attachment_id"])
    if attachment["size"] > store.limits.max_process_bytes:
        raise Fault("PROCESS_SIZE_LIMIT")
    store.room(identity, "cache", attachment["size"]+store.limits.script_output_bytes)
    return store.create_job(identity, params["attachment_id"], {"mode": "full"}, "script", params)


def execute_script(original, payload, out, limits, row):
    from .engines import version
    publication = out
    out = publication/"execution"
    out.mkdir(mode=0o700)
    if Path(original).stat().st_size > limits.max_process_bytes:
        raise Fault("PROCESS_SIZE_LIMIT")
    # Stable name provided to the script; immutable snapshot is never mounted writable by Landlock.
    input_folder = publication/"inputs" if limits.workspace_mode else out
    input_folder.mkdir(mode=0o700, exist_ok=True)
    incoming = input_folder/("input"+Path(original).suffix.lower())
    shutil.copyfile(original, incoming)
    before = digest_file(original)
    if incoming.suffix.lower() in (".xlsx", ".docx", ".ods"):
        import zipfile
        with zipfile.ZipFile(incoming) as archive:
            if sum(i.file_size for i in archive.infolist()) > limits.max_docx_expanded_bytes:
                raise Fault("EXPANDED_SIZE_LIMIT")
    script = out/"operation.py"
    script.write_text(payload["code"], encoding="utf-8")
    code_hash = digest_file(script)
    warnings = ["Temporary Python result; validation checks cover declared assertions, not all semantics.",
                "XLSX formulas are not recalculated; macros and workbook layout may not survive openpyxl edits."]
    if limits.script_isolation == "development":
        warnings.append("DEVELOPMENT MODE: no filesystem/network isolation; trusted local tests only.")
    launcher = Path(__file__).with_name("script_launcher.py")
    environment = {k: os.environ[k] for k in ("SYSTEMROOT", "WINDIR", "PATH") if k in os.environ}
    environment.update(TMPDIR=str(out), TMP=str(out), TEMP=str(out), OMP_NUM_THREADS=str(limits.engine_threads),
                       OPENBLAS_NUM_THREADS=str(limits.engine_threads), PYTHONNOUSERSITE="1", FILETOOLS_INPUT=str(incoming))
    started = time.monotonic()
    input_allowance = 0 if limits.workspace_mode else incoming.stat().st_size
    with (out/"stdout.log").open("wb") as stdout, (out/"stderr.log").open("wb") as stderr:
        from .process_tree import ManagedChild
        tree = ManagedChild([limits.script_python or sys.executable, "-I", str(launcher), str(out),
            limits.script_isolation, str(limits.worker_memory_mb*1024**2), str(limits.script_timeout_seconds),
            str(limits.script_output_bytes), str(incoming)] if limits.workspace_mode else [
            limits.script_python or sys.executable, "-I", str(launcher), str(out), limits.script_isolation,
            str(limits.worker_memory_mb*1024**2), str(limits.script_timeout_seconds), str(limits.script_output_bytes)],
            cwd=out, env=environment, stdout=stdout, stderr=stderr)
        child = tree.process
        try:
            while child.poll() is None:
                if time.monotonic()-started > limits.script_timeout_seconds:
                    raise Fault("SCRIPT_TIMEOUT")
                size = sum(p.stat().st_size for p in out.rglob("*") if p.is_file() and not p.is_symlink())
                if size > input_allowance+limits.script_output_bytes+2*1024**2:
                    raise Fault("SCRIPT_OUTPUT_LIMIT")
                if any((out/n).stat().st_size > 1024**2 for n in ("stdout.log", "stderr.log")):
                    raise Fault("SCRIPT_LOG_LIMIT")
                time.sleep(0.05)
        finally:
            tree.close()  # Always, including a normally exited leader with living descendants.
    if digest_file(original) != before:
        raise Fault("ORIGINAL_CHANGED")
    if digest_file(script) != code_hash:
        raise Fault("SCRIPT_CHANGED")
    if child.returncode == 73:
        raise Fault("SCRIPT_OUTPUT_LIMIT")
    if child.returncode == 74:
        raise Fault("SCRIPT_MEMORY_LIMIT")
    if sys.platform == "linux":
        import signal
        if child.returncode == -signal.SIGXFSZ:
            raise Fault("SCRIPT_OUTPUT_LIMIT")
        if child.returncode == -signal.SIGXCPU:
            raise Fault("SCRIPT_TIMEOUT")
    if child.returncode:
        raise Fault("SCRIPT_FAILED")
    if sum(p.stat().st_size for p in out.rglob("*") if p.is_file() and not p.is_symlink()) > input_allowance+limits.script_output_bytes+2*1024**2:
        raise Fault("SCRIPT_OUTPUT_LIMIT")
    assets, evidence = [], []
    for name in payload["outputs"]:
        file = out/name
        if not file.is_file() or file.is_symlink() or not 0 < file.stat().st_size <= limits.script_output_bytes:
            raise Fault("OUTPUT_NOT_VERIFIED")
        detected = identify(file, name)
        if name.lower().endswith(".xlsx"):
            import openpyxl
            with __import__("zipfile").ZipFile(file) as archive:
                if sum(i.file_size for i in archive.infolist()) > limits.max_docx_expanded_bytes:
                    raise Fault("EXPANDED_SIZE_LIMIT")
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=False)
            cached = openpyxl.load_workbook(file, read_only=True, data_only=True)
            try:
                for check in payload.get("checks", []):
                    if check["file"] != name:
                        continue
                    sheet = workbook[check["sheet"]]
                    if sheet.max_row < check.get("min_rows", 1) or sheet.max_column < check.get("min_columns", 1):
                        raise Fault("OUTPUT_ASSERTION_FAILED")
                    if "cell" in check and sheet[check["cell"]].value != check.get("equals"):
                        raise Fault("OUTPUT_ASSERTION_FAILED")
                for sheet in workbook:
                    for cells in sheet.iter_rows():
                        for cell in cells:
                            if cell.data_type == "f" and cached[sheet.title][cell.coordinate].value is None:
                                warnings.append(f"{name}/{sheet.title}/{cell.coordinate}: formula result unavailable; never assume zero.")
                                break
                        if len(warnings) >= 20:
                            break
                evidence.append({"filename": name, "worksheets": workbook.sheetnames})
            finally:
                workbook.close()
                cached.close()
        else:
            if any(c["file"] == name for c in payload.get("checks", [])):
                raise Fault("CHECK_FORMAT_UNSUPPORTED")
            if detected["category"] == "other":
                warnings.append(name+": only existence/size/hash validated; format semantics unverified.")
        assets.append({"artifact_id": uuid.uuid4().hex, "file": "execution/"+name, "kind": "output",
                       "bytes": file.stat().st_size, "sha256": digest_file(file), "actual_type": detected["actual_type"]})
    for name in ("operation.py", "stdout.log", "stderr.log"):
        assets.append({"artifact_id": uuid.uuid4().hex, "file": "execution/"+name, "kind": "script" if name.endswith(".py") else "log",
                       "bytes": (out/name).stat().st_size, "sha256": digest_file(out/name)})
    return {"status": "SUCCEEDED", "segments": [{"segment": 1, "text": json.dumps({"operation": payload["description"],
        "verified_outputs": evidence or payload["outputs"], "original_sha256": before, "script_sha256": code_hash}, ensure_ascii=False),
        "source": {"paragraph": 1}, "engine": "temporary-python"}], "assets": assets, "failures": [], "warnings": warnings,
        "config": {"mode": "full", "isolation": limits.script_isolation, "python": sys.version.split()[0]},
        "engines": {"openpyxl": version("openpyxl")}, "coverage": {"kind": "script", "full_document": False,
        "requested": payload["description"], "processed": [{"outputs": payload["outputs"]}], "checks": payload.get("checks", [])}}
