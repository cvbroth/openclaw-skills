"""Opt-in source-page pagination, verified against actual office rendering."""

from copy import deepcopy
import json
import re
from docx import Document
from docx.enum.text import WD_BREAK
import fitz
from ..document_sample import write_docx
from .word_styles import apply_styles
from .preview_worker import export_docx


def options(value):
    value = value or {"mode": "continuous"}
    if not isinstance(value, dict) or set(value) - {
        "mode",
        "font_fit",
        "minimum_font_pt",
        "allow_continuation",
    }:
        raise ValueError("pagination: 不支持的设置")
    result = {
        "mode": "continuous",
        "font_fit": False,
        "minimum_font_pt": 9,
        "allow_continuation": False,
        **value,
    }
    if result["mode"] not in {"continuous", "source-pages"}:
        raise ValueError("pagination.mode: 请选择连续排版或保留原件分页")
    if any(type(result[k]) is not bool for k in ["font_fit", "allow_continuation"]):
        raise ValueError("pagination: 开关必须是布尔值")
    v = result["minimum_font_pt"]
    if type(v) not in {int, float} or not 8 <= v <= 16:
        raise ValueError("pagination.minimum_font_pt: 字号下限范围 8–16 pt")
    return result


def group_pages(pages):
    groups = []
    for page in pages:
        n = page.get("source_pdf_physical_page")
        if type(n) is not int or n < 1:
            raise ValueError("SOURCE_PAGE_MAPPING_REQUIRED: 文字稿缺少可靠的原件物理页映射")
        if groups and n < groups[-1]["source_page"]:
            raise ValueError("SOURCE_PAGE_ORDER_CONFLICT: 原件映射逆序，请先核对")
        if groups and groups[-1]["source_page"] == n:
            groups[-1]["items"].extend(page["items"])
        else:
            groups.append({"source_page": n, "items": list(page["items"])})
    if not groups:
        raise ValueError("SOURCE_PAGE_MAPPING_REQUIRED: 没有可排版片段")
    return groups


def body_text(page, page_numbers=True):
    # Generated footer consists solely of the document page counter. Do not
    # remove digits from body or normalize away source text differences.
    blocks = page.get_text("blocks")
    if page_numbers and blocks and re.fullmatch(r"\s*文档第\s*\d+\s*页\s*", blocks[-1][4]):
        blocks = blocks[:-1]
    return "".join(re.sub(r"\s", "", b[4]) for b in blocks)


def generate_source_pages(output, pages, template, font_path, policy):
    groups = group_pages(pages)
    report = {
        "schema": "source-pagination-v1",
        "policy": policy,
        "blank_source_pages": "retain one output page",
        "groups": [],
        "verified": False,
    }
    report_path = output / "pagination-report.json"
    chosen = []
    expected_text = []
    counter = 1
    for i, group in enumerate(groups):
        entry = {"source_page": group["source_page"], "attempts": [], "blank": not group["items"]}
        report["groups"].append(entry)
        attempts = [(False, 1.0)]
        if group["items"]:
            attempts.append((True, 1.0))
            if policy["font_fit"]:
                base = (
                    (template.get("word_styles") or {})
                    .get("paragraphs", {})
                    .get("Normal", {})
                    .get("size_pt", template["parameters"]["body_pt"])
                )
                if policy["minimum_font_pt"] > base:
                    raise ValueError("pagination.minimum_font_pt: 下限不能高于正文字号")
                attempts.append((True, policy["minimum_font_pt"] / base))
        for attempt, (compact, scale) in enumerate(attempts):
            directory = output / "pagination-work" / f"group-{i + 1}-attempt-{attempt + 1}"
            directory.mkdir(parents=True, exist_ok=True)
            path = directory / "group.docx"
            write_docx(path, group["items"] or [("Normal", "")], font_path, template["parameters"])
            apply_styles(
                path,
                template.get("word_styles"),
                compact=compact,
                font_scale=scale,
                min_pt=policy["minimum_font_pt"],
                materialize=True,
            )
            pdf = export_docx(path, directory, timeout=35)
            with fitz.open(pdf) as document:
                texts = [body_text(p, template["parameters"].get("page_numbers", True)) for p in document]
            entry["attempts"].append(
                {
                    "reduce_spacing": compact,
                    "font_scale": scale,
                    "minimum_font_pt": policy["minimum_font_pt"],
                    "actual_output_pages": len(texts),
                }
            )
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
            if len(texts) == 1:
                break
        if len(texts) > 1 and not policy["allow_continuation"]:
            entry["conflict"] = "spacing/font floor exhausted; explicit continuation required"
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
            raise ValueError(
                f"SOURCE_PAGE_OVERFLOW: 原件第{group['source_page']}页实际需要{len(texts)}个输出页；请调整模板或明确允许续页"
            )
        entry["output_pages"] = list(range(counter, counter + len(texts)))
        counter += len(texts)
        expected_text.extend(texts)
        chosen.append(path)
    merged = Document(chosen[0])
    for path in chosen[1:]:
        paragraph = merged.add_paragraph()
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        for child in Document(path).element.body:
            if child.tag.endswith("}sectPr"):
                continue
            merged.element.body.insert(len(merged.element.body) - 1, deepcopy(child))
    result = output / "document.docx"
    merged.save(result)
    pdf = export_docx(result, output, timeout=35)
    with fitz.open(pdf) as document:
        actual = [body_text(p, template["parameters"].get("page_numbers", True)) for p in document]
    if actual != expected_text:
        report["conflict"] = "final rendered page content differs from isolated source groups"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
        raise ValueError("SOURCE_PAGE_FINAL_MAPPING_CONFLICT: 合并后实际分页与来源组不一致")
    report.update(
        verified=True,
        verification="actual DOCX→PDF rendering; per-page body comparison with isolated groups; footer counter excluded",
    )
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    return report
