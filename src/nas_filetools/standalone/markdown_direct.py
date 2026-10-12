"""Reader-faithful Markdown export. No question parser, option packing or OCR."""

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import subprocess
import fitz
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from ..reading_style import snapshot
from ..document_fonts import font_audit, font_runs, apply_run_font, embed_fonts
from .content_structure import parse_markdown, items
from .preview_worker import export_docx


def direct_template():
    style = snapshot()
    raw = json.dumps(style, ensure_ascii=False, sort_keys=True).encode()
    return {
        "record": {
            "id": "markdown-direct",
            "version": style["version"],
            "name": "Markdown原文直出",
            "sha256": hashlib.sha256(raw).hexdigest(),
        },
        "reading_style": style,
        "font": {"path": style["font_path"], "sha256": style["font_sha256"]},
        "font_fallbacks": style["fallbacks"],
    }


def compact_options(value):
    value = value or {}
    if not isinstance(value, dict) or set(value) - {
        "enabled",
        "minimum_font_pt",
        "minimum_line_multiple",
        "minimum_paragraph_em",
    }:
        raise ValueError("紧凑适配设置无效")
    result = {
        "enabled": False,
        "minimum_font_pt": 10,
        "minimum_line_multiple": 1.3,
        "minimum_paragraph_em": 0.3,
        **value,
    }
    if type(result["enabled"]) is not bool:
        raise ValueError("紧凑适配须明确选择")
    for key, low, high in [
        ("minimum_font_pt", 8, 10.5),
        ("minimum_line_multiple", 1.2, 1.85),
        ("minimum_paragraph_em", 0, 1),
    ]:
        if type(result[key]) not in {int, float} or not low <= result[key] <= high:
            raise ValueError(f"紧凑适配 {key} 超出范围 {low}–{high}")
    return result


def check_views(artifact, total, selected=None):
    views = artifact.get("pages", [])
    if not views:
        raise ValueError("原文直出缺少来源页映射；请使用自定义模板连续排版，不能推断来源页。")
    previous = 0
    for v in views:
        n = v.get("source_page")
        known = set(v.get("source_pages", [])) | ({n} if type(n) is int else set())
        if selected is not None and known and not known.intersection(selected):
            continue
        if type(n) is not int or n < 1 or type(total) is not int or n > total:
            raise ValueError("原文直出缺少可靠的来源物理页映射；片段位置不是原件页码。")
        if not v.get("data"):
            raise ValueError(f"来源第{n}页缺少文字记录资源，不能当作成功空白页导出。")
        if "source_pages" in v and (
            not isinstance(v["source_pages"], list)
            or any(type(x) is not int for x in v["source_pages"])
            or set(v["source_pages"]) != {n}
        ):
            raise ValueError(f"来源第{n}页片段涉及多个页或映射不明确，不能精确拆分直出。")
        if n < previous:
            raise ValueError(f"来源第{n}页映射逆序，不能改变正文顺序。")
        previous = n


def write_reader_docx(path, models, style, body_pt, line, paragraph_em):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = [Mm(v) for v in style["paper_mm"]]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(
        style["margin_mm"]
    )
    for model in models:
        for block in model["blocks"]:
            level = int(block["style"][-1]) if block["type"] == "heading" else None
            size = body_pt * style["heading_scale"][level - 1] if level else body_pt
            p = doc.add_paragraph(style=f"Heading {level}" if level else "Normal")
            fmt = p.paragraph_format
            # CSS line-height is an absolute multiple of the declared font size.
            # Word's MULTIPLE uses font metrics and makes CJK lines much taller.
            fmt.line_spacing = Pt(size * line)
            gap = size * style["heading_margin_em"][level - 1] if level else size * paragraph_em
            fmt.space_before = fmt.space_after = Pt(gap)
            fmt.left_indent = fmt.right_indent = fmt.first_line_indent = Pt(0)
            fmt.keep_with_next = False  # Reader headings do not bind across physical source pages.
            fmt.keep_together = False
            fmt.widow_control = False
            if block.get("quote_depth"):
                fmt.left_indent = Pt(size * (block["quote_depth"] * style["quote_indent_em"] + 0.5))
            if block["style"] == "Code":
                fmt.left_indent = fmt.right_indent = Pt(size * style["code_indent_em"])
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F4F6F9")
                p._p.get_or_add_pPr().append(shade)
            if block["style"] == "Separator" or block.get("quote_depth"):
                borders = OxmlElement("w:pBdr")
                border = OxmlElement("w:bottom" if block["style"] == "Separator" else "w:left")
                for key, value in {"val": "single", "sz": "6", "color": "AAB6C7", "space": "4"}.items():
                    border.set(qn("w:" + key), value)
                borders.append(border)
                p._p.get_or_add_pPr().append(borders)
            for part in block["runs"]:
                for name, text in font_runs(part["text"], style["font_path"]):
                    run = p.add_run(text)
                    apply_run_font(run, name or style["font_family"])
                    run.font.size = Pt(size)
                    run.font.bold = bool(level or part.get("bold"))
                    run.font.italic = bool(part.get("italic"))
                    run.font.color.rgb = RGBColor.from_string("243145")
    if not doc.paragraphs:
        doc.add_paragraph()  # A real empty paragraph keeps a blank source page when merging.
    doc.save(path)
    embed_fonts(path)


def page_text(page):
    return re.sub(r"\s", "", page.get_text())


def adaptation_attempts(style, policy):
    values = [(style["body_pt"], style["line_multiple"], style["paragraph_em"])]
    if policy["enabled"]:
        values.append((style["body_pt"], policy["minimum_line_multiple"], policy["minimum_paragraph_em"]))
        size = style["body_pt"] - 0.5
        while size > policy["minimum_font_pt"]:
            values.append((size, policy["minimum_line_multiple"], policy["minimum_paragraph_em"]))
            size -= 0.5
        if policy["minimum_font_pt"] < style["body_pt"]:
            values.append(
                (policy["minimum_font_pt"], policy["minimum_line_multiple"], policy["minimum_paragraph_em"])
            )
    return list(dict.fromkeys(values))


def validate_pdf(path, expected, n, single=True):
    with fitz.open(path) as doc:
        if single and len(doc) != 1:
            return False, len(doc)
        if "".join(page_text(p) for p in doc) != re.sub(r"\s", "", expected):
            raise ValueError(f"来源第{n}页PDF正文或顺序不一致，未发布。")
        for page in doc:
            if any(
                code == 0xFFFD or (glyph == 0 and not chr(code).isspace())
                for span in page.get_texttrace()
                for code, glyph, _, _ in span["chars"]
            ):
                raise ValueError(f"来源第{n}页PDF有缺字，未发布。")
            if any(
                b[0] < -1 or b[1] < -1 or b[2] > page.rect.width + 1 or b[3] > page.rect.height + 1
                for b in page.get_text("blocks")
            ):
                raise ValueError(f"来源第{n}页有文字越界，未发布。")
        return len(doc) == 1, len(doc)


def generate(request):
    from .reading_pdf import ReaderPDF
    from .tasks import atomic

    out = Path(request["output"])
    style = request["template_data"]["reading_style"]
    if snapshot() != style:
        raise ValueError("阅读样式或字体与快照不一致，请建立新任务；历史文件保留。")
    policy = compact_options(request.get("direct_compact"))
    purpose = request.get("direct_purpose", "export")
    formats = request.get("formats", ["docx", "pdf"])
    if purpose == "preview":
        formats = ["pdf"]
    report = {
        "schema": "markdown-direct-pagination-v2",
        "policy": policy,
        "purpose": purpose,
        "blank_missing": request.get("blank_missing", False),
        "placeholder_candidates": request.get("placeholder_candidates", []),
        "placeholders": request.get("placeholders", []),
        "groups": [],
        "verified": False,
        "ready": False,
        "issues": deepcopy(request.get("source_issues", [])),
        "content_states": request.get("content_states", []),
        "formats": formats,
        "source_artifact_id": request["source_artifact_id"],
        "source_sha256": request["source_sha256"],
        "reading_style": style,
        "renderer": {},
        "requested_source_pages": request.get("requested_source_pages"),
        "excluded_source_pages": request.get("excluded_source_pages", []),
    }
    groups = []
    for page in request["pages"]:
        n = page.get("source_pdf_physical_page")
        if type(n) is not int:
            report["issues"].append({"source_page": n, "reason": "缺少来源物理页映射"})
            continue
        if not groups or groups[-1]["source_page"] != n:
            groups.append({"source_page": n, "pages": []})
        groups[-1]["pages"].append(page)
    for issue in report["issues"] + report["placeholders"]:
        n = issue["source_page"]
        if type(n) is int and not any(g["source_page"] == n for g in groups):
            groups.append({"source_page": n, "pages": []})
    groups.sort(key=lambda g: g["source_page"])
    placeholder_pages = {i["source_page"]: i["reason"] for i in report["placeholders"]}
    blocked = {i["source_page"] for i in report["issues"]}
    checked, words, pdfs, text_by_page, fonts = 0, [], [], [], []
    renderer = ReaderPDF(out, style) if "pdf" in formats else None

    def persist(current=None):
        atomic(out / "pagination-report.json", report)
        atomic(
            out / "progress.json",
            {
                "phase": "逐页检查与生成",
                "completed": checked,
                "total": len(groups),
                "source_page": current,
                "blocked": len(report["issues"]),
            },
        )

    try:
        for group in groups:
            n = group["source_page"]
            entry = {"source_page": n, "attempts": [], "status": "BLOCKED"}
            report["groups"].append(entry)
            persist(n)
            if n in blocked:
                entry["reason"] = "正文或来源证据不完整，详见检查问题"
                checked += 1
                persist(n)
                continue
            try:
                placeholder = n in placeholder_pages
                models = [parse_markdown(p["text"], p.get("source_locator")) for p in group["pages"]]
                if placeholder:
                    if any(p["text"].strip() for p in group["pages"]):
                        raise ValueError("空白占位不能抹掉已有正文")
                    models = []
                if not placeholder and any(not p["text"].strip() for p in group["pages"]):
                    raise ValueError("没有正文，不能作为成功提取的空白页导出")
                labels = {
                    "POSSIBLE_TABLE_PRESERVED_AS_TEXT": "真实表格（含表头分隔行）",
                    "UNSUPPORTED_BLOCK_PRESERVED_AS_TEXT": "引用块或分隔线",
                    "IMAGE_REFERENCE_NOT_RENDERED": "图片引用",
                    "LINK_DISPLAY_TEXT_ONLY": "超链接",
                    "INLINE_UNSUPPORTED": "特殊行内格式",
                    "HEADING_LEVEL_4_TO_6_RENDERED_AS_3": "四至六级标题",
                    "LITERAL_BLOCK_PRESERVED": "代码块或HTML块",
                }
                diagnostics = {
                    labels.get(d["code"], d["code"])
                    for m in models
                    for d in m["diagnostics"]
                    if not d.get("supported")
                }
                if diagnostics:
                    raise ValueError(
                        "目前原文直出不支持的格式：" + "、".join(sorted(diagnostics)) + "；原文保留"
                    )
                fonts.append(
                    font_audit(
                        [{"source_pdf_physical_page": n, "items": [i for m in models for i in items(m)]}],
                        style["font_path"],
                    )
                )
                content = "".join(b["text"] for m in models for b in m["blocks"])
                fits = False
                for attempt, (size, line, gap) in enumerate(adaptation_attempts(style, policy)):
                    folder = out / "pagination-work" / f"group-{n}-attempt-{attempt + 1}"
                    folder.mkdir(parents=True)
                    metrics = {"body_pt": size, "line_multiple": line, "paragraph_em": gap}
                    entry["attempts"].append(metrics)
                    fits = True
                    if renderer:
                        pdf, measurement = renderer.render(
                            models, folder, size, line, gap, purpose == "preview"
                        )
                        metrics["pdf"] = measurement
                        report["renderer"]["pdf"] = measurement["renderer"]
                        pdf_fits, count = validate_pdf(pdf, content, n, single=False)
                        metrics["actual_output_pages"] = count
                        fits = pdf_fits and not measurement["overflow"]
                    if "docx" in formats and fits:
                        word = folder / "group.docx"
                        write_reader_docx(word, models, style, size, line, gap)
                        if "".join(p.text for p in Document(word).paragraphs) != content:
                            raise ValueError("Word正文或顺序不一致")
                        word_pdf = export_docx(word, folder, timeout=60)
                        word_fits, count = validate_pdf(word_pdf, content, n, single=False)
                        metrics["word_output_pages"] = count
                        metrics["actual_output_pages"] = max(metrics.get("actual_output_pages", 1), count)
                        fits &= word_fits
                        report["renderer"]["docx"] = {
                            "name": "python-docx",
                            "version": "1.1.2",
                            "validation": subprocess.check_output(
                                ["libreoffice", "--version"], text=True, timeout=10
                            ).strip(),
                            "role": "Word validation only; final PDF uses Chromium",
                        }
                    persist(n)
                    if fits:
                        break
                if purpose == "preview" and renderer:
                    report["preview_images"] = [
                        str(f.relative_to(out))
                        for f in sorted(
                            folder.glob("preview-*.png"), key=lambda f: int(f.stem.split("-")[-1])
                        )
                    ]
                    report["preview_pdf"] = str(pdf.relative_to(out))
                    report["preview_parameters"] = metrics
                if not fits:
                    count = metrics.get("actual_output_pages", 2)
                    raise ValueError(
                        "SOURCE_PAGE_OVERFLOW: "
                        + (
                            f"需要{count}张输出页；"
                            if count > 1
                            else "候选虽为1页，但正文及段落边距高度超过可用纸张；"
                        )
                        + (
                            "紧凑适配已到所设下限仍无法容纳"
                            if policy["enabled"]
                            else "未启用紧凑适配，可明确选择边界后重试"
                        )
                    )
                entry.update(
                    status="PLACEHOLDER" if placeholder else "READY",
                    placeholder_reason=placeholder_pages.get(n),
                    output_pages=[len(text_by_page) + 1],
                    content_equal=True,
                    final_parameters={"body_pt": size, "line_multiple": line, "paragraph_em": gap},
                )
                text_by_page.append(re.sub(r"\s", "", content))
                if "docx" in formats:
                    words.append(word)
                if renderer:
                    pdfs.append(pdf)
            except (ValueError, RuntimeError, TimeoutError) as exc:
                reason = str(exc) or "本页浏览器排版超过60秒或打印超过30秒，停止本页生成"
                entry["reason"] = reason
                report["issues"].append(
                    {
                        "source_page": n,
                        "category": "overflow"
                        if "SOURCE_PAGE_OVERFLOW" in reason
                        else "format"
                        if "不支持的格式" in reason
                        else "render",
                        "placeholder_eligible": False,
                        "reason": reason,
                    }
                )
            checked += 1
            persist(n)
        report["ready"] = bool(groups) and not report["issues"]
        report["verified"] = report["ready"]
        persist()
        atomic(out / "font-resolution.json", {"pages": fonts, "text_changed": False})
        atomic(
            out / "content-structure.json",
            {
                "schema": "derived-content-v1",
                "mode": "markdown-direct",
                "source_artifact_id": request["source_artifact_id"],
                "source_sha256": request["source_sha256"],
                "reading_style": style,
                "pages": groups,
            },
        )
        atomic(out / "layout-validation.json", report)
        if purpose != "export":
            return
        if not report["ready"]:
            reasons = "；".join(f"来源第{i['source_page']}页：{i['reason']}" for i in report["issues"])
            raise ValueError("原文直出未发布。" + reasons)
        if words:
            merged = Document(words[0])
            for path in words[1:]:
                body = [deepcopy(c) for c in Document(path).element.body if c.tag != qn("w:sectPr")]
                props = body[0].find(qn("w:pPr"))
                if props is None:
                    props = OxmlElement("w:pPr")
                    body[0].insert(0, props)
                props.append(OxmlElement("w:pageBreakBefore"))
                for child in body:
                    merged.element.body.insert(len(merged.element.body) - 1, child)
            merged.save(out / "document.docx")
            embed_fonts(out / "document.docx")
            validation = out / "word-validation"
            validation.mkdir()
            word_pdf = export_docx(out / "document.docx", validation, timeout=max(60, 15 + len(words) * 4))
            with fitz.open(word_pdf) as doc:
                if [page_text(p) for p in doc] != text_by_page:
                    raise ValueError("合并后Word来源分页或正文发生变化，未发布")
        if pdfs:
            with fitz.open() as combined:
                for pdf in pdfs:
                    with fitz.open(pdf) as part:
                        combined.insert_pdf(part)
                combined.save(out / "document.pdf")
            with fitz.open(out / "document.pdf") as doc:
                if [page_text(p) for p in doc] != text_by_page:
                    raise ValueError("合并后PDF来源分页或正文发生变化，未发布")
    finally:
        if renderer:
            renderer.close()
