"""Pandoc owns document content; Python supplies references, resources and audits."""

from copy import deepcopy
from pathlib import Path
import json
import re
import shutil
import subprocess
from collections import Counter
from zipfile import ZipFile
import fitz
from lxml import etree
from ..artifact_project import digest, relative
from ..document_entry import FONT_PATH
from ..document_fonts import font_audit, font_runs, embed_fonts
from .table_structure import tables
from .preview_worker import export_docx
from .markdown_rules import preserve_question_numbers, preserve_list_starts

READER = "commonmark_x+pipe_tables+raw_html+hard_line_breaks-yaml_metadata_block-smart-fancy_lists-tex_math_dollars-implicit_figures-footnotes-definition_lists-bracketed_spans-attributes-fenced_divs-raw_attribute-subscript-superscript-strikeout-task_lists-emoji-implicit_header_references-alerts"


def reference_path():
    return Path(__file__).with_name("reference.docx")


def reference_snapshot(path):
    return {
        "id": "pandoc-reference",
        "version": "1.0.0",
        "sha256": digest(path),
        "body_pt": 10.5,
        "line_multiple": 1.35,
        "paper_mm": [210, 297],
        "margin_mm": 18,
        "font": "Droid Sans Fallback",
    }


def run(args, **kwargs):
    try:
        completed = subprocess.run(
            ["pandoc", *args],
            capture_output=True,
            text=True,
            check=True,
            timeout=kwargs.pop("timeout", 120),
            **kwargs,
        )
    except subprocess.CalledProcessError as exc:
        raise ValueError("Pandoc转换失败：" + (exc.stderr or str(exc))[:2000]) from exc
    if completed.stderr.strip():
        raise ValueError("Pandoc报告无法完整转换：" + completed.stderr.strip())
    return completed.stdout


def ast_text(node):
    if isinstance(node, list):
        return "".join(ast_text(n) for n in node)
    if not isinstance(node, dict):
        return ""
    t = node.get("t")
    c = node.get("c")
    if t == "Str":
        return c
    if t in {"Code", "Math"}:
        return c[-1]
    if t == "CodeBlock":
        return c[-1]
    if t == "Image":
        return ""  # Alt text is retained in image accessibility metadata, not visible body.
    if t in {"Space", "SoftBreak", "LineBreak"}:
        return " "
    if t == "Table":
        return ast_text(c[1]) + ast_text(c[3]) + ast_text(c[4]) + ast_text(c[5])
    return ast_text(c) if isinstance(c, (list, dict)) else ""


def native_table(model):
    from .table_structure import validate

    validate(model)
    head_count = 0
    for r in range(model["rows"]):
        cells = [c for c in model["cells"] if c["row"] == r]
        if cells and all(c.get("header") for c in cells):
            head_count = r + 1
        else:
            break
    rows = []
    for r in range(model["rows"]):
        cells = []
        for cell in sorted((c for c in model["cells"] if c["row"] == r), key=lambda c: c["column"]):
            if r < head_count and r + cell["rowspan"] > head_count:
                raise ValueError("表头合并跨越正文边界，请明确全部表头行")
            inline = []
            for run in cell.get("runs", [{"text": cell["text"]}]):
                for index, line in enumerate(run["text"].split("\n")):
                    if index:
                        inline.append({"t": "LineBreak"})
                    if not line:
                        continue
                    part = (
                        {"t": "Code", "c": [["", [], []], line]}
                        if run.get("code")
                        else {"t": "Str", "c": line}
                    )
                    if run.get("italic"):
                        part = {"t": "Emph", "c": [part]}
                    if run.get("bold"):
                        part = {"t": "Strong", "c": [part]}
                    if run.get("href"):
                        part = {"t": "Link", "c": [["", [], []], [part], [run["href"], ""]]}
                    inline.append(part)
            cells.append(
                [
                    ["", [], []],
                    {
                        "t": {"left": "AlignLeft", "center": "AlignCenter", "right": "AlignRight"}.get(
                            cell.get("alignment"), "AlignDefault"
                        )
                    },
                    cell["rowspan"],
                    cell["colspan"],
                    [{"t": "Plain", "c": inline}] if inline else [],
                ]
            )
        rows.append([["", [], []], cells])
    return {
        "t": "Table",
        "c": [
            ["", [], []],
            [None, []],
            [
                [{"t": "AlignDefault"}, {"t": "ColWidth", "c": 1 / model["columns"]}]
                for _ in range(model["columns"])
            ],
            [["", [], []], rows[:head_count]],
            [[["", [], []], 0, [], rows[head_count:]]],
            [["", [], []], []],
        ],
    }


def replace_raw(node):
    if isinstance(node, list):
        return [replace_raw(n) for n in node]
    if not isinstance(node, dict):
        return node
    node = deepcopy(node)
    if node.get("t") == "RawBlock" and node["c"][0] == "html":
        raw = node["c"][1]
        if re.match(r"\s*<table(?:\s|>)", raw, re.I):
            from .table_structure import parse_html

            return native_table(parse_html(raw))
        if re.fullmatch(r"\s*<!-- source physical page \d+; task [\w-]+ -->\s*", raw):
            return {"t": "Null"}
        return {"t": "CodeBlock", "c": [["", [], []], raw]}
    if (
        node.get("t") == "RawInline"
        and node["c"][0] == "html"
        and re.fullmatch(r"<br\s*/?>", node["c"][1], re.I)
    ):
        return {"t": "LineBreak"}
    if node.get("t") == "RawInline":
        return {"t": "Code", "c": [["", [], []], node["c"][1]]}
    if "c" in node:
        node["c"] = replace_raw(node["c"])
    if "blocks" in node:
        node["blocks"] = replace_raw(node["blocks"])
    return node


def resolve_images(node, root, source, out):
    if isinstance(node, list):
        for n in node:
            resolve_images(n, root, source, out)
    elif isinstance(node, dict):
        if node.get("t") == "Link":
            url = node["c"][-1][0]
            if re.match(r"^[a-zA-Z][\w+.-]*:", url) and not re.match(r"^(https?|mailto):", url, re.I):
                raise ValueError("不支持的链接协议：" + url)
            resolve_images(node["c"][1], root, source, out)
        elif node.get("t") == "Image":
            url = node["c"][-1][0]
            if re.match(r"^[a-zA-Z][\w+.-]*:", url) or url.startswith(("/", "\\")):
                raise ValueError("图片资源必须为项目内本地文件，禁止远端自动下载：" + url)
            path = (root / source).parent / url
            try:
                path = relative(root, str(path.relative_to(root)))
            except (ValueError, OSError):
                raise ValueError("图片资源越过项目数据根：" + url) from None
            if not path.is_file():
                raise ValueError("图片资源缺失：" + url)
            if path.stat().st_size > 20_000_000:
                raise ValueError("图片资源过大：" + url)
            from PIL import Image

            with Image.open(path) as image:
                width, height = image.size
                image.verify()
            scale = min(174 / (width * 25.4 / 96), 230 / (height * 25.4 / 96), 1)
            node["c"][0][2].extend(
                [
                    ["width", f"{width * 25.4 / 96 * scale:.3f}mm"],
                    ["height", f"{height * 25.4 / 96 * scale:.3f}mm"],
                ]
            )
            resources = out / "resources"
            resources.mkdir(exist_ok=True)
            dest = resources / (digest(path) + path.suffix)
            shutil.copyfile(path, dest)
            node["c"][-1][0] = str(dest)
        else:
            resolve_images(node.get("c"), root, source, out)
            if "blocks" in node:
                resolve_images(node["blocks"], root, source, out)


def font_postprocess(path):
    # Split only textual OOXML runs when supplementary fonts are needed. Never rebuild paragraphs/tables.
    from docx.oxml.ns import qn

    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    for name, data in list(parts.items()):
        if not re.fullmatch(r"word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml", name):
            continue
        tree = etree.fromstring(data)
        for r in list(tree.iter(qn("w:r"))):
            texts = list(r.findall(qn("w:t")))
            if len(texts) != 1:
                continue
            text = texts[0].text or ""
            segments = font_runs(text, FONT_PATH)
            if not any(font for font, _ in segments):
                continue
            if any(child.tag not in {qn("w:rPr"), qn("w:t")} for child in r):
                raise ValueError("含特殊域的补充字体不能安全拆分")
            parent = r.getparent()
            index = parent.index(r)
            for font, part in segments:
                new = deepcopy(r)
                new.find(qn("w:t")).text = part
                new.find(qn("w:t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                props = new.find(qn("w:rPr"))
                if props is None:
                    props = etree.Element(qn("w:rPr"))
                    new.insert(0, props)
                fonts = props.find(qn("w:rFonts"))
                if fonts is None:
                    fonts = etree.SubElement(props, qn("w:rFonts"))
                for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
                    fonts.set(qn("w:" + key), font or "Droid Sans Fallback")
                parent.insert(index, new)
                index += 1
            parent.remove(r)
        parts[name] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    from zipfile import ZIP_DEFLATED

    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    embed_fonts(path)


def word_list_markers(path):
    """Audit actual OOXML automatic list labels, distinct from literal exam numbering."""
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path) as z:
        doc = etree.fromstring(z.read("word/document.xml"))
        numbering = etree.fromstring(z.read("word/numbering.xml"))

    def q(n):
        return "{" + ns["w"] + "}" + n

    abstracts = {x.get(q("abstractNumId")): x for x in numbering.findall(q("abstractNum"))}
    nums = {x.get(q("numId")): x for x in numbering.findall(q("num"))}
    counts = {}
    labels = []
    for paragraph in doc.findall(".//" + q("p")):
        props = paragraph.find(q("pPr") + "/" + q("numPr"))
        if props is None:
            continue
        nid = props.find(q("numId")).get(q("val"))
        level = int(props.find(q("ilvl")).get(q("val")))
        num = nums[nid]
        abstract = abstracts[num.find(q("abstractNumId")).get(q("val"))]
        lvl = next(x for x in abstract.findall(q("lvl")) if int(x.get(q("ilvl"))) == level)
        start_node = lvl.find(q("start"))
        start = int(start_node.get(q("val"))) if start_node is not None else 1
        override = next((x for x in num.findall(q("lvlOverride")) if int(x.get(q("ilvl"))) == level), None)
        if override is not None and override.find(q("startOverride")) is not None:
            start = int(override.find(q("startOverride")).get(q("val")))
        key = (nid, level)
        counts[key] = counts.get(key, start - 1) + 1
        pattern = lvl.find(q("lvlText")).get(q("val"))
        fmt = lvl.find(q("numFmt")).get(q("val"))
        if fmt == "bullet":
            labels.append(pattern)
            continue

        def value(match):
            v = counts.get((nid, int(match.group(1)) - 1), start)
            if fmt == "decimal":
                return str(v)
            if fmt in {"lowerLetter", "upperLetter"} and 1 <= v <= 26:
                return chr((97 if fmt == "lowerLetter" else 65) + v - 1)
            if fmt in {"lowerRoman", "upperRoman"}:
                result = ""
                left = v
                for n, symbol in [
                    (1000, "M"),
                    (900, "CM"),
                    (500, "D"),
                    (400, "CD"),
                    (100, "C"),
                    (90, "XC"),
                    (50, "L"),
                    (40, "XL"),
                    (10, "X"),
                    (9, "IX"),
                    (5, "V"),
                    (4, "IV"),
                    (1, "I"),
                ]:
                    while left >= n:
                        result += symbol
                        left -= n
                return result.lower() if fmt == "lowerRoman" else result
            raise ValueError("不支持验证的自动列表格式：" + fmt)

        labels.append(re.sub(r"%([1-9])", value, pattern))
    return "".join(labels)


def audit_tables(ast, xml):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    expected = []

    def visit(value):
        if isinstance(value, list):
            for item in value:
                visit(item)
        elif isinstance(value, dict):
            if value.get("t") == "Table":
                expected.append(value["c"])
            else:
                for item in value.values():
                    visit(item)

    visit(ast)
    actual = xml.xpath("//w:tbl", namespaces=ns)
    if len(actual) != len(expected):
        raise ValueError("Word可编辑表格数量与原生结构不一致")
    report = []
    for index, (native, word) in enumerate(zip(expected, actual), 1):
        rows = native[3][1] + [r for body in native[4] for r in body[2] + body[3]] + native[5][1]
        cells = [c for row in rows for c in row[1]]
        vertical = sum(c[2] for c in cells if c[2] > 1)
        spans = Counter(str(c[3]) for c in cells if c[3] > 1 for _ in range(c[2]))
        actual_spans = Counter(v for v in word.xpath(".//w:gridSpan/@w:val", namespaces=ns) if v != "1")
        actual_vertical = len(word.xpath(".//w:vMerge", namespaces=ns))
        if (
            actual_vertical != vertical
            or actual_spans != spans
            or len(word.xpath("./w:tr", namespaces=ns)) != len(rows)
        ):
            raise ValueError(f"表格{index}的Word合并关系或行数与结构不一致")
        report.append(
            {
                "table": index,
                "rows": len(rows),
                "columns": len(native[2]),
                "vertical_merge_cells": vertical,
                "horizontal_spans": dict(spans),
                "editable": True,
            }
        )
    return report


def prepare_reference(source, dest, title, settings, template=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(source)
    for section in doc.sections:
        section.header.paragraphs[0].text = title if settings["new_header"] else ""
        footer = section.footer.paragraphs[0]
        footer.clear()
        if settings["new_page_numbers"]:
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1"
            run.append(text)
            field.append(run)
            footer._p.append(field)
    if template:
        # Legacy style configuration affects reference styles only; no old generator is invoked.
        from docx.shared import Pt, Mm

        params = template["parameters"]
        for key, style in [
            ("body_pt", "Normal"),
            ("h1_pt", "Heading 1"),
            ("h2_pt", "Heading 2"),
            ("h3_pt", "Heading 3"),
        ]:
            if key in params:
                doc.styles[style].font.size = Pt(params[key])
        for section in doc.sections:
            for key in ["top", "bottom", "left", "right"]:
                value = params.get("margin_" + key + "_mm")
                if value is not None:
                    setattr(section, key + "_margin", Mm(value))
    doc.save(dest)
    if template and template.get("word_styles"):
        from .word_styles import apply_styles

        apply_styles(dest, template["word_styles"], materialize=False)


def generate(request):
    out = Path(request["output"])
    out.mkdir(parents=True, exist_ok=True)
    phase = "prepare"

    def progress(value, completed=0, total=1):
        nonlocal phase
        phase = value
        from .tasks import atomic

        atomic(out / "progress.json", {"phase": value, "completed": completed, "total": total})

    try:
        progress("parse")
        reference = Path(request["reference"])
        if digest(reference) != request["reference_sha256"]:
            raise ValueError("reference.docx快照校验失败")
        root = Path(request["project_root"])
        groups = []
        mapping = []
        body = []
        api = None
        for i, page in enumerate(request["pages"]):
            text = preserve_question_numbers(page["text"])
            records = tables(text)
            native = {}
            for table in records:
                if table.get("review_required"):
                    raise ValueError(
                        f"来源{page.get('source_pdf_physical_page')}表格待复核：{table['review_required']}"
                    )
                if table.get("error"):
                    raise ValueError(
                        f"来源{page.get('source_pdf_physical_page')}表格待复核：{table['error']}"
                    )
                if table["original"].lstrip().lower().startswith("<table"):
                    token = "ft-table-" + table["id"]
                    native[token] = native_table(table["model"])
            for table in reversed(records):
                token = "ft-table-" + table["id"]
                if token in native:
                    text = text[: table["start"]] + "\n\n<!-- " + token + " -->\n\n" + text[table["end"] :]
            ast = json.loads(run(["-f", READER, "-t", "json"], input=text))
            api = ast["pandoc-api-version"]

            def replace_tables(value):
                if isinstance(value, list):
                    return [replace_tables(v) for v in value]
                if not isinstance(value, dict):
                    return value
                if value.get("t") == "RawBlock":
                    match = re.fullmatch(r"\s*<!-- (ft-table-[a-f0-9]+) -->\s*", value["c"][1])
                    if match:
                        return native[match.group(1)]
                return {k: replace_tables(v) for k, v in value.items()}

            ast = preserve_list_starts(replace_raw(replace_tables(ast)), text)
            resolve_images(
                ast,
                root,
                str(Path(request["resource_base"]) / "content.md")
                if request.get("resource_base")
                else request["source_path"],
                out,
            )
            bookmark = "ft-source-" + str(i + 1)
            body.append({"t": "Div", "c": [[bookmark, [], []], ast["blocks"]]})
            groups.append(
                {
                    "source_pdf_physical_page": page.get("source_pdf_physical_page"),
                    "source_locator": page["source_locator"],
                    "items": [("Normal", ast_text(ast["blocks"]))],
                }
            )
            mapping.append(
                {
                    "source_page": page.get("source_pdf_physical_page"),
                    "source_locator": page["source_locator"],
                    "bookmark": bookmark,
                    "output_pages": None,
                    "pagination": "continuous; source boundaries do not determine output pages",
                }
            )
            progress("parse", i + 1, len(request["pages"]))
        if not body:
            raise ValueError("没有可导出的正文")
        document = {"pandoc-api-version": api, "meta": {}, "blocks": body}
        (out / "pandoc.json").write_text(json.dumps(document, ensure_ascii=False))
        progress("fonts")
        fonts = font_audit(groups, FONT_PATH)
        (out / "font-resolution.json").write_text(json.dumps(fonts, ensure_ascii=False, indent=2))
        prepare_reference(
            reference,
            out / "reference-used.docx",
            request["title"],
            request["settings"],
            request.get("template_data"),
        )
        progress("pandoc-docx")
        run(
            [
                "-f",
                "json",
                "-t",
                "docx",
                "--reference-doc=" + str(out / "reference-used.docx"),
                "--output=" + str(out / "document.docx"),
                str(out / "pandoc.json"),
            ],
            timeout=min(7200, max(180, len(request["pages"]) * 15)),
        )
        font_postprocess(out / "document.docx")
        progress("validate-word")
        with ZipFile(out / "document.docx") as z:
            xml = etree.fromstring(z.read("word/document.xml"))
            actual = "".join(xml.itertext())  # Includes OOXML property text only in w:t, filter below.
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            actual = "".join(xml.xpath("//w:t/text()", namespaces=ns))

        def canonical(s):
            return re.sub(r"\s", "", s)

        table_checks = audit_tables(document, xml)
        expected = ast_text(document["blocks"])
        if canonical(actual) != canonical(expected):
            (out / "content-diff.json").write_text(
                json.dumps({"expected": expected, "actual": actual}, ensure_ascii=False)
            )
            raise ValueError("Pandoc生成Word正文与原生AST文字/顺序不一致，未发布文件")
        progress("libreoffice-pdf")
        pdf = export_docx(out / "document.docx", out, timeout=min(7200, max(180, len(request["pages"]) * 20)))
        progress("validate-pdf")
        with fitz.open(pdf) as doc:
            page_texts = [
                canonical(p.get_text(clip=fitz.Rect(0, 45, p.rect.width, p.rect.height - 45))) for p in doc
            ]
            expected_inventory = Counter(canonical(expected + word_list_markers(out / "document.docx")))
            actual_inventory = Counter("".join(page_texts))
            header_signatures = []

            def headers(value):
                if isinstance(value, list):
                    for item in value:
                        headers(item)
                elif isinstance(value, dict):
                    if value.get("t") == "Table":
                        signature = canonical(ast_text(value["c"][3]))
                        if signature:
                            header_signatures.append(signature)
                    else:
                        for item in value.values():
                            headers(item)

            headers(document)
            repeated_headers = []
            # A complete registered table header may repeat at the top of a continuation page.
            # Account for its exact excess inventory; never discard body text in the PDF.
            for index, text in enumerate(page_texts):
                if index == 0:
                    continue
                for signature in header_signatures:
                    counts = Counter(signature)
                    if text.startswith(signature) and all(
                        actual_inventory[k] - expected_inventory[k] >= v for k, v in counts.items()
                    ):
                        actual_inventory.subtract(counts)
                        repeated_headers.append({"output_page": index + 1, "table_header": signature})
                        break
            if +actual_inventory != expected_inventory or any(v < 0 for v in actual_inventory.values()):
                (out / "pdf-content-diff.json").write_text(
                    json.dumps(
                        {
                            "missing": dict(expected_inventory - actual_inventory),
                            "extra": dict(actual_inventory - expected_inventory),
                        },
                        ensure_ascii=False,
                    )
                )
                raise ValueError("PDF正文字符与Word不一致（含裁切/丢字/额外文字），未发布文件")
            missing = [
                i + 1
                for i, p in enumerate(doc)
                if any(
                    code == 0xFFFD or (glyph == 0 and not chr(code).isspace())
                    for span in p.get_texttrace()
                    for code, glyph, origin, bbox in span["chars"]
                )
            ]
            outside = [
                i + 1
                for i, p in enumerate(doc)
                if any(
                    b[0] < -1 or b[1] < -1 or b[2] > p.rect.width + 1 or b[3] > p.rect.height + 1
                    for b in p.get_text("blocks")
                )
            ]
            if missing or outside:
                raise ValueError(f"PDF字体/边界验证失败：缺字页{missing}，越界页{outside}")
            count = len(doc)
        renderer = {
            "pandoc": run(["--version"]).splitlines()[0],
            "libreoffice": subprocess.check_output(
                ["libreoffice", "--version"], text=True, timeout=10
            ).strip(),
            "chain": "Pandoc native AST -> DOCX -> LibreOffice PDF",
            "reference_sha256": digest(out / "reference-used.docx"),
        }
        (out / "content-structure.json").write_text(
            json.dumps(
                {
                    "schema": "pandoc-derived-v1",
                    "source_artifact_id": request["source_artifact_id"],
                    "source_sha256": request["source_sha256"],
                    "source_mapping": mapping,
                    "ast": document,
                    "registered_metadata": request.get("metadata", {}),
                },
                ensure_ascii=False,
            )
        )
        (out / "layout-validation.json").write_text(
            json.dumps(
                {
                    "content_equal": True,
                    "tables": table_checks,
                    "pdf_character_inventory_equal": True,
                    "repeated_table_headers": repeated_headers,
                    "output_pages": count,
                    "source_mapping": mapping,
                    "renderer": renderer,
                    "cleaning_summary": request["settings"]["cleaning_summary"],
                    "visual_review": "requires actual rendered-file review; automated checks insufficient",
                    "forced_source_breaks": False,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
        progress("finished", 1, 1)
    except Exception as exc:
        raise ExportError(phase, str(exc)) from exc


class ExportError(ValueError):
    def __init__(self, stage, message):
        super().__init__(message)
        self.stage = stage
