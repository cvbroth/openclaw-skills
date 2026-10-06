"""Conservative OCR question sample structuring and fixed DOCX/PDF layout.

Standalone on purpose: the existing restricted Python tool can submit this source
as code and consume its sole registered input via FILETOOLS_INPUT.
"""
import hashlib
import html
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_FAMILY = "Droid Sans Fallback"
PAGE_MARK = re.compile(r"^##\s*(?:第\s*(\d+)\s*页|Page\s+(\d+))\s*$", re.I)
QUESTION = re.compile(r"^(\d{1,3})[.．、]\s*(.+)$")
OPTION = re.compile(r"^([A-DＡ-Ｄ])[.．、:：)）]\s*(.*)$")
CHAPTER = re.compile(r"^(?:#{1,3}\s+(?!第\s*\d+\s*页)(.+)|第[一二三四五六七八九十百\d]+[章节篇]\s*.*)$")


def join_scan_lines(parts):
    """Join within one semantic block only; retain Latin word boundaries."""
    result = ""
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if result and result[-1].isascii() and result[-1].isalnum() and part[0].isascii() and part[0].isalnum():
            result += " "
        result += part
    return result


def structure_ocr(raw, pages=(6, 10), evidence=None, reviewed_edits=None):
    """Return independent blocks plus untouched line references and audit records.

    Repeated boundary headers require three distinct source pages. Position
    evidence may additionally name exact (page,line,text) margin fragments.
    Ambiguous numbering/inline options remain verbatim and receive an issue.
    """
    records, page = [], None
    for number, line in enumerate(raw.splitlines(), 1):
        mark = PAGE_MARK.match(line.strip())
        if mark:
            page = int(mark.group(1) or mark.group(2))
            continue
        if page is None and line.strip() and not line.strip().startswith("<!--"):
            raise ValueError("OCR_PAGE_MARKERS_REQUIRED")
        if page is not None and pages[0] <= page <= pages[1] and line.strip() and not line.strip().startswith("<!--"):
            records.append({"page": page, "line": number, "raw": line, "text": line.strip()})
    if {r["page"] for r in records} != set(range(pages[0], pages[1] + 1)):
        raise ValueError("OCR_SAMPLE_PAGES_MISSING")
    edit_map = {}
    for edit in reviewed_edits or []:
        key = (edit["page"], edit["line"])
        if key in edit_map or edit.get("verified") is not True or not edit.get("reference") or not edit.get("reason"):
            raise ValueError("INVALID_REVIEW_EVIDENCE")
        edit_map[key] = edit
    expanded = []
    applied = []
    for record in records:
        edit = edit_map.pop((record["page"], record["line"]), None)
        if not edit:
            expanded.append(record)
            continue
        if edit["original"] != record["raw"] or not isinstance(edit["replacement"], list) or not edit["replacement"]:
            raise ValueError("REVIEW_SOURCE_MISMATCH")
        for index, text in enumerate(edit["replacement"]):
            if not isinstance(text, str) or not text.strip() or "\n" in text:
                raise ValueError("INVALID_REVIEW_REPLACEMENT")
            expanded.append({**record, "text": text, "part": index, "review": edit})
        applied.append(edit)
    if edit_map:
        raise ValueError("REVIEW_SOURCE_NOT_FOUND")
    records = expanded
    boundaries = {}
    for source_page in range(pages[0], pages[1] + 1):
        lines = [r for r in records if r["page"] == source_page]
        for index in {0, len(lines) - 1}:
            boundaries.setdefault(lines[index]["text"], set()).add(source_page)
    margin_evidence = {}
    source_keys = {(r["page"], r["line"], r["text"]) for r in records}
    for item in evidence or []:
        key = (item["page"], item["line"], item["text"])
        if (item.get("region") not in ("header", "footer") or item.get("verified") is not True
                or not item.get("reference") or key in margin_evidence):
            raise ValueError("INVALID_MARGIN_EVIDENCE")
        if key not in source_keys:
            raise ValueError("MARGIN_SOURCE_MISMATCH")
        margin_evidence[key] = item
    blocks, issues, removed = [], [], []
    current, active, chapter = None, None, None

    def issue(code, fragments, detail):
        issues.append({"code": code, "detail": detail, "fragments": fragments})

    def finish():
        if current is None:
            return
        labels = [b["label"] for b in current["parts"] if b["kind"] == "option"]
        if labels != list("ABCD"):
            issue("OPTIONS_INCOMPLETE_OR_ORDER", current["fragments"], f"Observed options: {labels}; no options supplied.")
        stem = current["parts"][0]
        if stem["text"] and (stem["text"].count("“") > stem["text"].count("”") or stem["text"].endswith("著")):
            issue("POSSIBLE_TRUNCATION", stem["fragments"], "Unclosed quotation or suspicious final fragment; compare original page, do not complete it.")

    for record in records:
        text = record["text"]
        peers = [r for r in records if r["page"] == record["page"]]
        at_boundary = record is peers[0] or record is peers[-1]
        repeat = at_boundary and len(boundaries.get(text, set())) >= 3
        # Never classify a chapter/question/option as a repeated running header.
        repeat = repeat and not (CHAPTER.match(text) or QUESTION.match(text) or OPTION.match(text))
        positional = margin_evidence.get((record["page"], record["line"], text))
        if repeat or positional:
            removed.append({"fragment": record, "reason": "verified-margin-position" if positional else "same-exact-text-at-page-boundary",
                            "evidence": positional or {"pages": sorted(boundaries[text])}})
            continue
        heading = CHAPTER.match(text)
        if heading:
            finish()
            chapter = re.sub(r"^#{1,3}\s*", "", text)
            blocks.append({"kind": "chapter", "text": chapter, "fragments": [record]})
            current, active = None, None
            continue
        match = QUESTION.match(text)
        if match and re.match(r"^\d+[.]\d{1,2}(?:\D|$)", text):
            match = None  # Unspaced decimal body numbers are not question anchors.
        if match:
            number = int(match.group(1))
            expected = current is None or number == current["number"] + 1
            observed = [b["label"] for b in current["parts"] if b["kind"] == "option"] if current else []
            completed = bool(observed) and observed[-1] == "D"
            if expected and (current is None or completed):
                finish()
                active = {"kind": "stem", "text": match.group(2), "fragments": [record]}
                current = {"kind": "question", "number": number, "chapter": chapter, "parts": [active], "fragments": [record]}
                blocks.append(current)
                continue
            issue("AMBIGUOUS_NUMBERING", [record], "Numbered line is not a proven new question; preserved inside the current block.")
        match = OPTION.match(text)
        if match and current:
            label = unicodedata.normalize("NFKC", match.group(1))
            previous = [b["label"] for b in current["parts"] if b["kind"] == "option"]
            if label in previous or (previous and ord(label) != ord(previous[-1]) + 1) or (not previous and label != "A"):
                issue("OPTION_ORDER", [record], f"Unexpected option {label}; kept in source order.")
            active = {"kind": "option", "label": label, "text": match.group(2), "fragments": [record]}
            current["parts"].append(active)
            current["fragments"].append(record)
            continue
        if current and text[0] in "①②③④":
            active = {"kind": "point", "text": text, "fragments": [record]}
            current["parts"].append(active)
            current["fragments"].append(record)
            continue
        if re.search(r"\s+[B-D][.．、]", text):
            issue("INLINE_OPTIONS_UNRESOLVED", [record], "Multiple inline option markers require original-page review; not split by guesswork.")
        if active is not None:
            active["text"] = join_scan_lines([active["text"], text])
            active["fragments"].append(record)
            current["fragments"].append(record)
        else:
            blocks.append({"kind": "unresolved", "text": text, "fragments": [record]})
            issue("UNRESOLVED_STRUCTURE", [record], "No established question/option context; text retained.")
    finish()
    preserved = [r for b in blocks for r in b["fragments"]] + [item["fragment"] for item in removed]
    if Counter((r["page"], r["line"], r["raw"]) for r in preserved) != Counter((r["page"], r["line"], r["raw"]) for r in records):
        raise ValueError("SOURCE_FRAGMENT_ACCOUNTING_FAILED")
    return {"schema_version": 1, "source_sha256": hashlib.sha256(raw.encode()).hexdigest(), "pages": list(pages),
            "blocks": blocks, "issues": issues, "removed_headers_footers": removed,
            "review_status": "original-page-comparison-required", "reviewed_edits": applied}


# One configuration controls both templates; no question/page-number exceptions.
DEFAULT_TEMPLATE = {
    "font_family": FONT_FAMILY, "body_pt": 11.5, "line_spacing": 1.35,
    "space_after_pt": 6, "margin_cm": 2, "width_cm": 21, "height_cm": 29.7,
    "title_pt": 16, "part_pt": 14, "chapter_pt": 13, "subject_pt": 13,
    "type_pt": 11.5, "source_pt": 9, "footer_pt": 9, "heading_before_pt": 3,
    "option_indent_cm": 0.65, "option_hanging_cm": 0.35,
    "short_question_max_fraction": 0.30, "short_option_max_lines": 3,
    "short_stem_max_lines": 3, "minimum_start_lines": 2,
    "footer_distance_cm": 1, "max_pages": 50,
}


def template_config(overrides=None):
    result = {**DEFAULT_TEMPLATE, **(overrides or {})}
    if set(result) != set(DEFAULT_TEMPLATE):
        raise ValueError("UNKNOWN_TEMPLATE_PARAMETER")
    if not 0 < result["short_question_max_fraction"] < 0.5:
        raise ValueError("INVALID_SHORT_QUESTION_THRESHOLD")
    if result["minimum_start_lines"] < 2:
        raise ValueError("INVALID_MINIMUM_START_LINES")
    return result


def source_label(fragments):
    pages = sorted({r["page"] for r in fragments})
    return f"原PDF：第{'、'.join(map(str, pages))}页"


def heading_style(text, previous=None):
    if re.match(r"^第.+部分", text):
        return "Heading 1"
    if re.match(r"^第.+[章节篇]", text) or text in ("导论", "绪论", "引言", "结语"):
        return "Heading 2"
    if "选择题" in text or re.match(r"^[一二三四五六七八九十]+、", text):
        return "Heading 3"
    return "Subject" if previous == "Heading 1" else "Heading 2"


def paragraphs(model, title):
    result = [("Title", title)]
    previous_heading = None
    for block in model["blocks"]:
        if block["kind"] == "chapter":
            style = heading_style(block["text"], previous_heading)
            result.append((style, block["text"]))
            previous_heading = style
        elif block["kind"] == "question":
            for part in block["parts"]:
                if part["kind"] == "stem":
                    text, style = f"{block['number']}.\u00a0{part['text']}", "Question"
                elif part["kind"] == "option":
                    text, style = f"{part['label']}. {part['text']}", "Option"
                else:
                    text, style = part["text"], "Point"
                result.append((style, text))
            result.append(("Source", source_label(block["fragments"])))
        else:
            result.extend([("Unresolved", "[结构待核实] " + block["text"]),
                           ("Source", source_label(block["fragments"]))])
    return result


def item_groups(items):
    """Bounded question groups, separated by Source; never bind all questions."""
    index = 0
    while index < len(items):
        start = index
        if items[index][0] == "Question":
            while index < len(items) and items[index][0] != "Source":
                index += 1
            if index == len(items):
                raise ValueError("QUESTION_SOURCE_REQUIRED")
        index += 1
        yield start, index, items[start:index]


class StoryLayout:
    """Actual Story measurements and drawing; no speculative break-* CSS."""
    def __init__(self, font_path, template=None):
        self.template = template_config(template)
        self.font_path = Path(font_path) if font_path else None
        t = self.template
        self.media = fitz.Rect(0, 0, t["width_cm"] * 72 / 2.54, t["height_cm"] * 72 / 2.54)
        inset = t["margin_cm"] * 72 / 2.54
        self.where = self.media + (inset, inset, -inset, -inset)
        self.line = t["body_pt"] * t["line_spacing"]
        self.short_limit = self.where.height * t["short_question_max_fraction"]
        self.cache = {}
        face = f'@font-face {{font-family: sample; src: url({self.font_path.name});}}' if self.font_path else ""
        self.css = face + f'''
        body {{font-family: sample; font-size:{t['body_pt']}pt; line-height:{t['line_spacing']}; margin:0; color:#000;}}
        p {{margin:0 0 {t['space_after_pt']}pt 0;}}
        .title {{font-size:{t['title_pt']}pt;}}
        .heading-1 {{font-size:{t['part_pt']}pt; font-weight:bold;}}
        .heading-2 {{font-size:{t['chapter_pt']}pt; font-weight:bold;}}
        .heading-3 {{font-size:{t['type_pt']}pt; font-weight:bold;}}
        .subject {{font-size:{t['subject_pt']}pt;}}
        .option,.point,.continuation {{margin-left:{t['option_indent_cm'] * 72 / 2.54}pt;}}
        .option,.point {{text-indent:{-t['option_hanging_cm'] * 72 / 2.54}pt;}}
        .source {{font-size:{t['source_pt']}pt;}}
        '''

    def story(self, items):
        content = "".join(f'<p class="{style.lower().replace(" ", "-")}">{html.escape(text)}</p>' for style, text in items)
        return fitz.Story("<html><body>" + content + "</body></html>", user_css=self.css,
                          archive=fitz.Archive(str(self.font_path.parent)) if self.font_path else None)

    def height(self, items):
        key = tuple(items)
        if key not in self.cache:
            more, filled = self.story(items).place(fitz.Rect(0, 0, self.where.width, 100000))
            if more:
                raise ValueError("MEASUREMENT_LIMIT")
            self.cache[key] = fitz.Rect(filled).height
        return self.cache[key]

    def short_paragraph(self, item, lines):
        return self.height([item]) <= lines * self.line + self.template["space_after_pt"] + 0.1

    def question_short(self, group):
        return self.height(group) <= self.short_limit

    def long_option_tail(self, item):
        """Split only presentation flow; keep at least two measured final lines.

        The single semantic option/text in the model and DOCX is unchanged.
        """
        style, text = item
        low, high = 1, len(text) - 1
        target = self.template["minimum_start_lines"] * self.line + self.template["space_after_pt"] - 0.1
        while low < high:
            size = (low + high) // 2
            if self.height([("Continuation", text[-size:])]) >= target:
                high = size
            else:
                low = size + 1
        cut = len(text) - low
        # Do not insert a semantic space or cut an ASCII word if avoidable.
        while cut > 0 and text[cut - 1].isascii() and text[cut - 1].isalnum() and text[cut].isascii() and text[cut].isalnum():
            cut -= 1
        if cut == 0:
            raise ValueError("UNBREAKABLE_LONG_OPTION_TAIL")
        return (style, text[:cut]), ("Continuation", text[cut:])


def _font_style(style, size, template):
    from docx.shared import RGBColor
    t = template
    style.font.name = t["font_family"]
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    color = style.element.rPr.find(qn("w:color"))
    for attribute in ("themeColor", "themeTint", "themeShade"):
        color.attrib.pop(qn("w:" + attribute), None)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), t["font_family"])
        fonts.attrib.pop(qn("w:" + key + "Theme"), None)
    style.paragraph_format.line_spacing = t["line_spacing"]
    style.paragraph_format.space_after = Pt(t["space_after_pt"])
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.widow_control = True
    borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if borders is not None:
        borders.getparent().remove(borders)


def write_docx(path, items, font_path=None, template=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    layout = StoryLayout(font_path, template)
    t = layout.template
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(t["width_cm"]), Cm(t["height_cm"])
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(t["margin_cm"])
    section.footer_distance = Cm(t["footer_distance_cm"])
    sizes = {"Normal": t["body_pt"], "Title": t["title_pt"], "Subject": t["subject_pt"],
             "Heading 1": t["part_pt"], "Heading 2": t["chapter_pt"], "Heading 3": t["type_pt"],
             "Question": t["body_pt"], "Option": t["body_pt"], "Point": t["body_pt"],
             "Source": t["source_pt"], "Unresolved": t["body_pt"], "Footer": t["footer_pt"]}
    for name, size in sizes.items():
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _font_style(style, size, t)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(t["heading_before_pt"])
        if name in ("Option", "Point"):
            style.paragraph_format.left_indent = Cm(t["option_indent_cm"])
            style.paragraph_format.first_line_indent = Cm(-t["option_hanging_cm"])
    added = []
    for style, text in items:
        paragraph = document.add_paragraph(text, style)
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = style in ("Title", "Subject", "Heading 1", "Heading 2", "Heading 3")
        added.append(paragraph)
    for start, end, group in item_groups(items):
        if group[0][0] != "Question":
            continue
        short = layout.question_short(group)
        for offset, item in enumerate(group):
            paragraph = added[start + offset]
            is_short_option = item[0] in ("Option", "Point") and layout.short_paragraph(item, t["short_option_max_lines"])
            paragraph.paragraph_format.keep_together = short or is_short_option
            paragraph.paragraph_format.keep_with_next = short and offset < len(group) - 1
        if not short:
            if layout.short_paragraph(group[0], t["short_stem_max_lines"]):
                added[start].paragraph_format.keep_together = True
                added[start].paragraph_format.keep_with_next = True
            # Bind only the final option's ending and its source, not the next question.
            added[end - 2].paragraph_format.keep_with_next = True
            added[end - 1].paragraph_format.keep_together = True
    footer = section.footer.paragraphs[0]
    footer.style = document.styles["Footer"]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("文档第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "1"
    run.append(value)
    field.append(run)
    footer._p.append(field)
    footer.add_run(" 页")
    document.core_properties.title = items[0][1]
    document.save(path)


def write_pdf(path, items, font_path, template=None):
    layout = StoryLayout(font_path, template)
    t = layout.template
    if not Path(font_path).is_file():
        raise ValueError("FONT_FILE_UNAVAILABLE")
    font = fitz.Font(fontfile=str(font_path))
    missing = sorted({c for _, text in items + [("Footer", "文档第 0123456789 页")] for c in text if not c.isspace() and not font.has_glyph(ord(c))})
    if missing:
        raise ValueError("FONT_MISSING_GLYPHS: " + "".join(missing))
    writer = fitz.DocumentWriter(str(path))
    device, y, page_count = None, layout.where.y0, 0
    decisions = []

    def new_page():
        nonlocal device, y, page_count
        if device is not None:
            writer.end_page()
        page_count += 1
        if page_count > t["max_pages"]:
            raise ValueError("SAMPLE_LAYOUT_PAGE_LIMIT")
        device = writer.begin_page(layout.media)
        y = layout.where.y0

    def draw_unit(unit, keep=False, minimum=0):
        nonlocal y
        height = layout.height(unit)
        if device is None:
            new_page()
        remaining = layout.where.y1 - y
        if y > layout.where.y0 and ((keep and height > remaining - 0.1) or remaining < minimum):
            new_page()
        story = layout.story(unit)
        while True:
            more, filled = story.place(fitz.Rect(layout.where.x0, y, layout.where.x1, layout.where.y1))
            story.draw(device)
            y = fitz.Rect(filled).y1
            if not more:
                break
            new_page()

    try:
        groups = list(item_groups(items))
        for group_index, (start, end, group) in enumerate(groups):
            if group[0][0] != "Question":
                reserve = 0
                if group[0][0] in ("Title", "Subject", "Heading 1", "Heading 2", "Heading 3"):
                    for _, _, following in groups[group_index + 1:]:
                        if following[0][0] != "Question":
                            reserve += layout.height(following)
                            continue
                        if layout.question_short(following):
                            reserve += layout.height(following)
                        else:
                            reserve += t["minimum_start_lines"] * layout.line + t["space_after_pt"]
                        break
                draw_unit(group, keep=True, minimum=layout.height(group) + reserve)
                continue
            short = layout.question_short(group)
            decisions.append({"start_item": start, "end_item": end, "short": short, "measured_pt": layout.height(group)})
            if short:
                draw_unit(group, keep=True)
                continue
            prefix = group[:-2]
            index = 0
            while index < len(prefix):
                item = prefix[index]
                unit = [item]
                keep = item[0] in ("Option", "Point") and layout.short_paragraph(item, t["short_option_max_lines"])
                if index == 0 and len(prefix) > 1 and layout.short_paragraph(item, t["short_stem_max_lines"]):
                    unit.append(prefix[1])
                    # A short stem stays with the next semantic part only if bounded.
                    if layout.height(unit) <= layout.short_limit:
                        keep = True
                        index += 1
                    else:
                        unit.pop()
                minimum = t["minimum_start_lines"] * layout.line + t["space_after_pt"]
                if index == 0 and len(prefix) > 1 and len(unit) == 1 and layout.short_paragraph(item, t["short_stem_max_lines"]):
                    minimum += layout.height(unit)  # Reserve two lines of a long first option.
                draw_unit(unit, keep=keep, minimum=minimum)
                index += 1
            last, source = group[-2:]
            if layout.short_paragraph(last, t["short_option_max_lines"]):
                draw_unit([last, source], keep=True)
            else:
                head, tail = layout.long_option_tail(last)
                draw_unit([head], minimum=t["minimum_start_lines"] * layout.line + t["space_after_pt"])
                draw_unit([tail, source], keep=True)
    finally:
        if device is not None:
            writer.end_page()
        writer.close()
    # Mature Page text API adds the document's own page number in the footer band.
    with fitz.open(path) as pdf:
        for number, page in enumerate(pdf, 1):
            label = f"文档第 {number} 页"
            width = font.text_length(label, fontsize=t["footer_pt"])
            page.insert_text(((page.rect.width - width) / 2, page.rect.height - t["footer_distance_cm"] * 72 / 2.54),
                             label, fontname="sample-footer", fontfile=str(font_path), fontsize=t["footer_pt"], color=(0, 0, 0))
        # Finalize font resources once after every Story and footer has drawn.
        # Native MuPDF subsetting requires no fontTools runtime dependency.
        # Full rewrite compresses streams and collects/merges duplicate objects;
        # an incremental save would retain the full duplicate footer font.
        pdf.subset_fonts()
        finalized = Path(path).with_suffix(".final.pdf")
        try:
            pdf.save(finalized, garbage=4, deflate=True)
        except BaseException:
            finalized.unlink(missing_ok=True)
            raise
    finalized.replace(path)
    return decisions


def canonical(text):
    return re.sub(r"\s+", "", text)


def page_body(page, number, template=None):
    t = template_config(template)
    bottom = page.rect.height - t["margin_cm"] * 72 / 2.54
    footer = page.get_text(clip=fitz.Rect(0, bottom + 1, page.rect.width, page.rect.height), sort=True)
    if canonical(footer) != canonical(f"文档第 {number} 页"):
        raise ValueError("DOCUMENT_FOOTER_MISMATCH")
    return page.get_text(clip=fitz.Rect(0, 0, page.rect.width, bottom + 1), sort=True)


def validate_pair(docx_path, pdf_path, items, template=None):
    expected = [text for _, text in items]
    document = Document(docx_path)
    if [p.text for p in document.paragraphs] != expected:
        raise ValueError("DOCX_CONTENT_MISMATCH")
    with fitz.open(pdf_path) as pdf:
        actual = "".join(page_body(page, number, template) for number, page in enumerate(pdf, 1))
        if canonical(actual) != canonical("".join(expected)):
            raise ValueError("PDF_CONTENT_OR_ORDER_MISMATCH")
        geometry = []
        for number, page in enumerate(pdf, 1):
            traces = page.get_texttrace()
            if not traces or not page_body(page, number, template).strip():
                raise ValueError("EMPTY_PDF_PAGE")
            for span in traces:
                if not page.rect.contains(fitz.Rect(span["bbox"])):
                    raise ValueError("PDF_TEXT_CROPPED")
                for codepoint, glyph, origin, bbox in span["chars"]:
                    if codepoint == 0xfffd or (glyph == 0 and not chr(codepoint).isspace()):
                        raise ValueError("PDF_MISSING_GLYPH")
            geometry.append({"page": number, "words": len(page.get_text("words"))})
    return {"content_equal": True, "pdf_pages": geometry, "docx_paragraphs": len(document.paragraphs),
            "document_page_numbers": "checked", "visual_review": "required", "word_render": "not-checked-by-this-function"}


def generate_sample(input_path, output_dir, font_path, title="合成验证样本（非用户原册）", pages=(6, 10), template=None):
    """Generate a small sample only; no automatic full-book processing or saving."""
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    input_bytes = Path(input_path).read_bytes()
    if Path(input_path).suffix.lower() == ".json":
        review = json.loads(input_bytes)
        raw = review["raw_ocr"]
        raw_bytes = raw.encode("utf-8")
        model = structure_ocr(raw, pages, review.get("margin_evidence"), review.get("reviewed_edits"))
        model["review_notes"] = review.get("review_notes", [])
    else:
        raw_bytes = input_bytes
        raw = raw_bytes.decode("utf-8")
        model = structure_ocr(raw, pages)
    items = paragraphs(model, title)
    write_docx(output / "sample.docx", items, font_path, template)
    decisions = write_pdf(output / "sample.pdf", items, font_path, template)
    checks = validate_pair(output / "sample.docx", output / "sample.pdf", items, template)
    (output / "structured.json").write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "original-ocr.md").write_bytes(raw_bytes)
    lines = ["# 疑点及处理记录", "", "原始OCR不改写。下列修订仅依据提供的原页核对记录；生成器不自行验证记录，未确定处不得凭常识补字。", ""]
    for issue in model["issues"]:
        lines += [f"## {issue['code']}", issue["detail"]]
        for fragment in issue["fragments"]:
            lines.append(f"- 原PDF第{fragment['page']}页 / OCR行{fragment['line']}：{fragment['raw']}")
        lines.append("")
    lines += ["## 对照原页的独立修订（不写回OCR）", ""]
    for edit in model["reviewed_edits"]:
        lines += [f"- 第{edit['page']}页 / OCR行{edit['line']}；原文：{edit['original']}",
                  f"  修订：{' / '.join(edit['replacement'])}；原因：{edit['reason']}；依据：{edit['reference']}"]
    for note in model.get("review_notes", []):
        lines.append("- " + note)
    lines += ["## 移出正文的页眉页脚", ""]
    for removed in model["removed_headers_footers"]:
        r = removed["fragment"]
        lines.append(f"- 第{r['page']}页 / OCR行{r['line']}：{r['raw']}；证据：{json.dumps(removed['evidence'], ensure_ascii=False)}")
    (output / "issues.md").write_text("\n".join(lines), encoding="utf-8")
    checks.update(template=template_config(template), pdf_layout_decisions=decisions, word_font_embedding="not-embedded; desktop substitution remains unverified", source_sha256=model["source_sha256"], issues=len(model["issues"]), font=FONT_FAMILY,
                  font_sha256=hashlib.sha256(Path(font_path).read_bytes()).hexdigest(), original_pdf_review="supplied-manual-evidence" if model["reviewed_edits"] else "not-performed",
                  generator_verified_review_evidence=False)
    (output / "validation.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    return checks
