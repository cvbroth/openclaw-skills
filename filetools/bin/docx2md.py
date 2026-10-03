#!/usr/bin/env python3
"""Word 转 Markdown。优先用 pandoc（效果最好），没装就用 python-docx 简易转换。

用法:
    python3 docx2md.py 文档.docx [-o 文档.md]
"""
import argparse
import shutil
import subprocess


def via_pandoc(inp, out):
    subprocess.run(["pandoc", inp, "-o", out], check=True)


def via_pythondocx(inp, out):
    from docx import Document
    doc = Document(inp)
    lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = p.style.name
        if style.startswith("Heading 1"):
            lines.append(f"# {t}")
        elif style.startswith("Heading 2"):
            lines.append(f"## {t}")
        elif style.startswith("Heading 3"):
            lines.append(f"### {t}")
        else:
            lines.append(t)
        lines.append("")
    for table in doc.tables:
        for row in table.rows:
            lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        lines.append("")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main():
    ap = argparse.ArgumentParser(description="Word 转 Markdown")
    ap.add_argument("input", help="docx 路径")
    ap.add_argument("-o", "--output", default=None)
    args = ap.parse_args()
    out = args.output or args.input.rsplit(".", 1)[0] + ".md"
    if shutil.which("pandoc"):
        via_pandoc(args.input, out)
    else:
        via_pythondocx(args.input, out)
    print(f"完成: {out}")


if __name__ == "__main__":
    main()
