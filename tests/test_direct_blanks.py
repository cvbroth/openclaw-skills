"""Actual blank-page pagination, scoped failures and newly shared basic formats."""

import json
import fitz
import pytest
from docx import Document
from nas_filetools.artifact_project import digest, page_data
from nas_filetools.standalone.direct_source import collect
from nas_filetools.standalone.content_structure import parse_markdown, reading_html
from test_template_exports import project as template_project, terminal


@pytest.fixture(name="project")
def fixture(tmp_path):
    yield from template_project.__wrapped__(tmp_path)


def set_text(m, pid, aid, index, text, failed=False):
    p = m.get(pid)
    a = next(a for a in p["artifacts"] if a["artifact_id"] == aid)
    v = a["pages"][index]
    root = m.path(pid)
    page_data(root, v["data"], {"text": text})
    v["data_sha256"] = digest(root / v["data"])
    if failed:
        v["error"] = True
    texts = [
        json.loads((root / x["data"]).read_text().split("=", 1)[1].rstrip(";\n"))["text"] for x in a["pages"]
    ]
    (root / a["path"]).write_text("\n\n".join(texts))
    a["sha256"] = digest(root / a["path"])
    m.save(root, p)
    return a


@pytest.mark.parametrize("formats", [["pdf"], ["docx"], ["pdf", "docx"]])
def test_real_blank_pages_and_immutable_source(project, formats):
    m, pid, aid = project
    root = m.path(pid)
    a = set_text(m, pid, aid, 4, "", True)
    original = (root / a["path"]).read_bytes()
    request = {"source_artifact_id": aid, "pages": "1-10", "formats": formats}
    with pytest.raises(ValueError, match="来源第1页"):
        m.export_documents(pid, request)
    summary = m.direct_source_summary(pid, request)
    assert summary["expected_output_pages"] == 10
    assert [i["source_page"] for i in summary["placeholder_candidates"]] == [1, 2, 3, 4, 5, 9]
    job = m.export_documents(pid, {**request, "blank_missing": True})
    p, t = terminal(m, pid, job["task_id"])
    assert t["status"] == "SUCCEEDED", t.get("error")
    outputs = [a for a in p["artifacts"] if a.get("task_id") == t["task_id"] and a.get("role") == "export"]
    assert {a["format"] for a in outputs} == set(formats)
    report = t["direct_report"]
    assert len(report["groups"]) == 10 and report["ready"]
    assert [g["source_page"] for g in report["groups"]] == list(range(1, 11))
    assert [g["output_pages"] for g in report["groups"]] == [[n] for n in range(1, 11)]
    assert [g["source_page"] for g in report["groups"] if g["status"] == "PLACEHOLDER"] == [1, 2, 3, 4, 5, 9]
    for a in outputs:
        assert len(a["export_source"]["placeholders"]) == 6
        assert a["source_output_mapping"] == report["groups"]
        rendered = (
            root / a["path"]
            if a["format"] == "pdf"
            else root / "diagnostics/format" / t["task_id"] / "work/word-validation/document.pdf"
        )
        with fitz.open(rendered) as pdf:
            assert len(pdf) == 10
            assert [i + 1 for i, p in enumerate(pdf) if not p.get_text().strip()] == [1, 2, 3, 4, 5, 9]
            assert (
                "第6页" in "".join(pdf[5].get_text().split())
                and "第二片段" in pdf[6].get_text()
                and "第10页" in "".join(pdf[9].get_text().split())
            )
    assert (
        root / next(a for a in p["artifacts"] if a["artifact_id"] == aid)["path"]
    ).read_bytes() == original
    assert next(a for a in p["artifacts"] if a["artifact_id"] == aid)["pages"][4]["error"]


def test_scope_partial_and_resources_never_blank(project):
    m, pid, aid = project
    root = m.path(pid)
    a = set_text(m, pid, aid, 4, "", True)
    assert not collect(root, m.get(pid), a, [8])["issues"]
    assert not collect(root, m.get(pid), a, [6, 7, 8])["issues"]
    a = set_text(m, pid, aid, 2, "", True)  # Empty fragment alongside nonempty fragment from page7.
    result = collect(root, m.get(pid), a, [7], True)
    assert result["issues"][0]["category"] == "partial_body" and not result["placeholders"]
    (root / a["pages"][3]["data"]).write_text("corrupt")
    result = collect(root, m.get(pid), a, [8], True)
    assert result["issues"] and all(not i["placeholder_eligible"] for i in result["issues"])
    with pytest.raises(ValueError, match="资源"):
        m.export_documents(
            pid, {"source_artifact_id": aid, "pages": "8", "formats": ["pdf"], "blank_missing": True}
        )


def test_new_basic_formats_real_word_pdf(project):
    m, pid, aid = project
    text = "#### 四级\n\n##### 五级\n\n###### 六级\n\n> 引用 **粗体** 与 *斜体*\n>\n> 第二段引用\n\n---\n\n```text\n代码 **文字** | 管道\n  缩进 <script>保持原样</script>\n```\n\n普通 `代码文字` ①⸺。"
    model = parse_markdown(text)
    html = reading_html(model)
    assert (
        "<h4>" in html
        and "<h5>" in html
        and "<h6>" in html
        and "md-quote" in html
        and "md-code" in html
        and "md-separator" in html
    )
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert not [d for d in model["diagnostics"] if not d.get("supported")]
    assert next(b for b in model["blocks"] if b["type"] == "separator")["text"] == "---"
    set_text(m, pid, aid, 0, text)
    job = m.export_documents(pid, {"source_artifact_id": aid, "pages": "6", "formats": ["docx", "pdf"]})
    p, t = terminal(m, pid, job["task_id"])
    assert t["status"] == "SUCCEEDED", t.get("error")
    for a in p["artifacts"]:
        if a.get("task_id") != t["task_id"] or a.get("role") != "export":
            continue
        if a["format"] == "docx":
            doc = Document(m.path(pid) / a["path"])
            body = "".join(p.text for p in doc.paragraphs)
            assert [p.style.name for p in doc.paragraphs[:3]] == ["Heading 4", "Heading 5", "Heading 6"]
        else:
            with fitz.open(m.path(pid) / a["path"]) as pdf:
                body = "".join(p.get_text() for p in pdf)
        for token in [
            "四级",
            "五级",
            "六级",
            "第二段引用",
            "---",
            "代码 **文字** | 管道",
            "<script>保持原样</script>",
            "①⸺",
        ]:
            assert token in body


@pytest.mark.parametrize(
    "body,reason", [("长段落。" * 5000, "OVERFLOW"), ("| A | B |\n| -- | -- |\n| C | D |", "不支持")]
)
def test_blank_policy_cannot_hide_existing_problem(project, body, reason):
    m, pid, aid = project
    set_text(m, pid, aid, 0, body)
    job = m.export_documents(
        pid, {"source_artifact_id": aid, "pages": "6", "formats": ["pdf"], "blank_missing": True}
    )
    p, t = terminal(m, pid, job["task_id"])
    assert t["status"] == "FAILED" and reason in t["error"]["message"]
    assert not [a for a in p["artifacts"] if a.get("task_id") == t["task_id"] and a.get("role") == "export"]
    assert not t["direct_report"]["placeholders"]


def test_cancel_queued_check_immediately_without_spawning(project):
    from unittest.mock import patch
    import time

    m, pid, aid = project
    original = m.run_child
    calls = []

    def observe(*args, **kwargs):
        calls.append(kwargs.get("module"))
        return original(*args, **kwargs)

    with patch.object(m, "run_child", side_effect=observe):
        with m.lock:
            job = m.export_documents(
                pid, {"source_artifact_id": aid, "pages": "8", "formats": ["pdf"], "purpose": "check"}
            )
            m.cancel(pid, job["task_id"])
            assert m.task(m.get(pid), job["task_id"])["status"] == "CANCELLED"
        deadline = time.monotonic() + 10
        while m.queue.unfinished_tasks and time.monotonic() < deadline:
            time.sleep(0.02)
        assert not m.queue.unfinished_tasks
    assert not calls


def test_multiple_empty_fragments_share_one_placeholder(project):
    m, pid, aid = project
    set_text(m, pid, aid, 1, "", True)
    a = set_text(m, pid, aid, 2, "", True)
    resolved = collect(m.path(pid), m.get(pid), a, [7], True)
    assert not resolved["issues"]
    assert len(resolved["placeholders"]) == 1
    assert len(resolved["placeholders"][0]["fragment_locators"]) == 2
    job = m.export_documents(
        pid, {"source_artifact_id": aid, "pages": "7", "formats": ["pdf"], "blank_missing": True}
    )
    p, t = terminal(m, pid, job["task_id"])
    assert t["status"] == "SUCCEEDED", t.get("error")
    a = next(a for a in p["artifacts"] if a.get("task_id") == job["task_id"] and a.get("role") == "export")
    with fitz.open(m.path(pid) / a["path"]) as pdf:
        assert len(pdf) == 1 and not pdf[0].get_text().strip()
