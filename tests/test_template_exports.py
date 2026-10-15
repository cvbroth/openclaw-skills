"""Template exports use immutable text and local office rendering, never recognition."""

import copy
import time
import uuid
from unittest.mock import patch

import fitz
import pytest
from docx import Document

from nas_filetools.artifact_project import digest, emit_viewer, page_data, register_artifact
from nas_filetools.standalone.document_exports import export_options
from nas_filetools.standalone.tasks import Projects
from test_standalone_http import configuration

TEMPLATE = {"id": "questions-zh-cn", "version": "1.0.0"}


@pytest.fixture
def project(tmp_path):
    config = configuration()
    config["conversion"]["generate_documents"] = True  # Old installations may retain this default.
    m = Projects(tmp_path / "store", config)
    try:
        path = tmp_path / "synthetic169.pdf"
        with fitz.open() as doc:
            for _ in range(169):
                doc.new_page()
            doc.save(path)
        pid = m.upload(path, path.name)["project_id"]
        # The local upload analysis also writes metadata. Finish it before tests
        # inject immutable synthetic extraction records outside the UI.
        deadline = time.monotonic() + 100
        while m.get(pid).get("analysis", {}).get("status") in {"QUEUED", "RUNNING"}:
            assert time.monotonic() < deadline, "source analysis did not settle"
            time.sleep(0.05)
        with m.lock:
            p = m.get(pid)
            root = m.path(pid)
            fragments = [(n, f"来源第{n}页完整文字。") for n in range(6, 11)]
            fragments.insert(2, (7, "来源第7页的第二片段必须保留。"))
            views = []
            for i, (n, text) in enumerate(fragments):
                resource = f"content/fragment-{i}.js"
                page_data(root, resource, {"text": text})
                views.append({"source_page": n, "locator": {"kind": "block", "id": str(i)}, "data": resource})
            (root / "content/selected.md").write_text("\n\n".join(t for _, t in fragments))
            a = register_artifact(
                root, p, name="合成文字稿", format="markdown", path="content/selected.md", pages=views
            )
            emit_viewer(root, p)
            m.save(root, p)
        yield m, pid, a["artifact_id"]
    finally:
        m.close()


def terminal(m, pid, tid):
    deadline = time.monotonic() + 100
    while time.monotonic() < deadline:
        p = m.get(pid)
        t = m.task(p, tid)
        if t["status"] not in {"QUEUED", "RUNNING", "CANCELLING"}:
            return p, t
        time.sleep(0.05)
    pytest.fail("bounded export timeout")


def export(m, pid, aid, **settings):
    return m.export_documents(pid, {"source_artifact_id": aid, "template": TEMPLATE, **settings})


@pytest.mark.parametrize("formats,pages", [(["docx"], "全部"), (["pdf"], "7-9"), (["docx", "pdf"], "7-9")])
def test_real_formats_ranges_fragments_and_provenance(project, formats, pages):
    m, pid, aid = project
    before = m.get(pid)
    job = export(m, pid, aid, formats=formats, pages=pages)
    p, t = terminal(m, pid, job["task_id"])
    assert t["status"] == "SUCCEEDED", t.get("error")
    outputs = [a for a in p["artifacts"] if a.get("task_id") == t["task_id"] and a.get("role") == "export"]
    assert {a["format"] for a in outputs} == set(formats)
    with fitz.open(m.path(pid) / "diagnostics/format" / t["task_id"] / "work/document.pdf") as pdf:
        text = "".join("".join(page.get_text().split()) for page in pdf)
        assert "第二片段必须保留" in text
        assert ("第6页" in text) == (pages == "全部")
        assert len(pdf) < len(t["document_export"]["included_source_pages"])
    for a in outputs:
        assert (
            a["export_source"]["sha256"]
            == next(x for x in before["artifacts"] if x["artifact_id"] == aid)["sha256"]
        )
        assert a["export_source"]["renderer"]["chain"] == "Pandoc native AST -> DOCX -> LibreOffice PDF"
        assert all(x["output_pages"] is None for x in a["source_output_mapping"])


def test_invalid_range_mapping_and_assembly(project):
    m, pid, aid = project
    p = m.get(pid)
    a = next(a for a in p["artifacts"] if a["artifact_id"] == aid)
    with pytest.raises(ValueError, match="没有文字记录"):
        export(m, pid, aid, pages="11")
    for value in ["0", "170", "9-7", "", [True]]:
        with pytest.raises(ValueError):
            export_options({"pages": value}, p, a)
    for formats in [[], ["pdf", "pdf"], ["html"], "pdf"]:
        with pytest.raises(ValueError):
            export_options({"formats": formats}, p, a)
    bad = copy.deepcopy(a)
    bad["pages"][0].pop("source_page")
    assert export_options({"pages": "全部"}, p, bad)["source_pages"] is None
    with pytest.raises(ValueError, match="映射"):
        export_options({"pages": "7"}, p, bad)
    with m.lock:
        p = m.get(pid)
        next(a for a in p["artifacts"] if a["artifact_id"] == aid)["pages"].pop()
        m.save(m.path(pid), p)
    with pytest.raises(ValueError, match="整稿与逐页映射不一致"):
        export(m, pid, aid)


def test_revision_history_repeat_failure_cancel_preserve_files(project):
    m, pid, aid = project
    info = m.text_artifact(pid, aid)
    new = m.save_text_revision(
        pid,
        {
            "request_id": str(uuid.uuid4()),
            "parent_artifact_id": aid,
            "parent_sha256": info["sha256"],
            "head_id": info["head_id"],
            "head_sha256": info["head_sha256"],
            "locator": info["pages"][1]["locator"],
            "text": "修订版第7页正文。",
        },
    )["artifact_id"]
    exported = []
    for source in [new, aid, new]:
        job = export(m, pid, source, formats=["docx"], pages="7")
        p, t = terminal(m, pid, job["task_id"])
        assert t["status"] == "SUCCEEDED", t.get("error")
        a = next(a for a in p["artifacts"] if a.get("task_id") == t["task_id"] and a.get("role") == "export")
        body = "".join(p.text for p in Document(m.path(pid) / a["path"]).paragraphs)
        assert ("修订版第7页正文" in body) == (source == new)
        exported.append(a)
    assert len({a["path"] for a in exported}) == 3

    # Simulate a completed child racing cancellation: cancellation still prevents publication.
    def cancel_result(request, *args, **kwargs):
        tid = next(t["task_id"] for t in m.get(pid)["tasks"] if t["status"] == "RUNNING")
        m.cancel(pid, tid)
        return {"status": "SUCCEEDED"}

    for response in [
        {"status": "FAILED", "error": {"code": "TEST", "message": "synthetic failure"}},
        cancel_result,
    ]:
        with patch.object(
            m, "run_child", side_effect=response if callable(response) else None, return_value=response
        ):
            job = export(m, pid, aid)
            p, t = terminal(m, pid, job["task_id"])
            assert t["status"] in {"FAILED", "CANCELLED"}
            assert not any(a.get("task_id") == t["task_id"] for a in p["artifacts"])
    for a in exported:
        assert digest(m.path(pid) / a["path"]) == a["sha256"]


def test_extraction_defaults_and_explicit_legacy_option(project):
    m, pid, _ = project
    for parameters in [None, {}, {"dpi": 150}, {"generate_documents": False}]:
        item = m.enqueue(pid, "fixture", [6], parameters)
        _, task = terminal(m, pid, item["task_id"])
        assert task["parameters"]["generate_documents"] is False
    with pytest.raises(ValueError, match="自动文档生成已迁移"):
        m.enqueue(pid, "fixture", [6], {"generate_documents": True})


def test_legacy_unmapped_full_text_allowed_range_refused(project):
    m, pid, aid = project
    with m.lock:
        p = m.get(pid)
        artifact = next(a for a in p["artifacts"] if a["artifact_id"] == aid)
        artifact["pages"] = []
        m.save(m.path(pid), p)
    with pytest.raises(ValueError, match="映射"):
        export(m, pid, aid, pages="7")
    job = export(m, pid, aid, formats=["docx"], pages="全部")
    p, task = terminal(m, pid, job["task_id"])
    assert task["status"] == "SUCCEEDED", task.get("error")
    a = next(a for a in p["artifacts"] if a.get("task_id") == task["task_id"] and a.get("role") == "export")
    body = "".join(p.text for p in Document(m.path(pid) / a["path"]).paragraphs)
    assert "第6页完整文字" in body and "第10页完整文字" in body and "第二片段必须保留" in body
    assert a["export_source"]["source_pages"] is None
    assert a["export_source"]["source_page_mapping_complete"] is False


def test_atomic_concurrent_writers_keep_independent_staging(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    import json
    import threading
    from nas_filetools.standalone.tasks import atomic

    target = tmp_path / "project.json"
    barrier = threading.Barrier(2)
    original = json.dump

    def simultaneous_dump(*args, **kwargs):
        original(*args, **kwargs)
        barrier.wait(timeout=5)

    with patch("nas_filetools.standalone.tasks.json.dump", side_effect=simultaneous_dump):
        with ThreadPoolExecutor(2) as pool:
            futures = [pool.submit(atomic, target, {"writer": i}) for i in range(2)]
            for future in futures:
                future.result()
    assert json.loads(target.read_text())["writer"] in {0, 1}
    assert not list(tmp_path.glob("*.next"))
