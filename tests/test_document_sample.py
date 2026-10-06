"""Synthetic semantics and real mature-library document checks; no user PDF."""
import hashlib
import json
import zipfile
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.oxml.ns import qn

from nas_filetools.document_sample import generate_sample, join_scan_lines, structure_ocr

FIXTURE = Path(__file__).parent / "fixtures" / "document_sample" / "ocr.md"


def test_merge_boundaries_and_page_provenance():
    model = structure_ocr(FIXTURE.read_text())
    questions = [b for b in model["blocks"] if b["kind"] == "question"]
    assert [b["number"] for b in questions] == [4, 5, 6, 7, 8, 9, 1, 2]
    assert questions[0]["parts"][1]["text"] == "保留原始材料，并记录资料来源。"
    cross = questions[5]
    assert cross["parts"][0]["text"] == "跨原PDF页的题干在整理稿中可以继续合并，但不按原页强制分页吗？"
    assert {r["page"] for r in cross["fragments"]} == {8, 9}
    assert [part["kind"] for part in questions[1]["parts"]] == ["stem"] + ["point"] * 4 + ["option"] * 4
    assert any(i["code"] == "POSSIBLE_TRUNCATION" and i["fragments"][0]["page"] == 6 for i in model["issues"])
    assert any(i["code"] == "OPTIONS_INCOMPLETE_OR_ORDER" for i in model["issues"])
    assert len(model["removed_headers_footers"]) == 10


def test_numbered_body_and_mid_option_header_remain_uncertain():
    raw = FIXTURE.read_text().replace("并记录资料来源。", "3. 正文编号不是已确认的新题。\n并记录资料来源。")
    raw = raw.replace("B. 直接改写不清楚的文字。", "B. 直接改写不清楚的文字。\n马原·单选")
    model = structure_ocr(raw)
    first = next(b for b in model["blocks"] if b["kind"] == "question")
    assert "3. 正文编号" in first["parts"][1]["text"]
    assert "马原·单选" in first["parts"][2]["text"]
    assert any(i["code"] == "AMBIGUOUS_NUMBERING" for i in model["issues"])
    assert len(model["removed_headers_footers"]) == 10


def test_missing_page_and_unmarked_text_fail_closed():
    with pytest.raises(ValueError, match="OCR_PAGE_MARKERS_REQUIRED"):
        structure_ocr("no source page")
    with pytest.raises(ValueError, match="OCR_SAMPLE_PAGES_MISSING"):
        structure_ocr("## 第6页\n1. unknown")


def test_latin_words_and_decimal_are_not_destroyed():
    assert join_scan_lines(["hello", "world"]) == "hello world"
    model = structure_ocr(FIXTURE.read_text().replace("内部的内容？", "3.5是正文数字。"))
    assert len([b for b in model["blocks"] if b["kind"] == "question"]) == 8


def test_real_docx_pdf_content_styles_and_exact_original(tmp_path):
    font = tmp_path / "font.ttf"
    font.write_bytes(fitz.Font("cjk").buffer)
    original = tmp_path / "ocr.md"
    original.write_bytes(FIXTURE.read_bytes().replace(b"\n", b"\r\n"))
    before = hashlib.sha256(original.read_bytes()).hexdigest()
    checks = generate_sample(original, tmp_path / "out", font)
    assert checks["content_equal"]
    assert (tmp_path / "out" / "original-ocr.md").read_bytes() == original.read_bytes()
    assert checks["source_sha256"] == before
    assert hashlib.sha256(original.read_bytes()).hexdigest() == before
    word = Document(tmp_path / "out" / "sample.docx")
    section = word.sections[0]
    assert abs(section.page_width.cm - 21) < 0.01 and abs(section.page_height.cm - 29.7) < 0.01
    assert all(abs(value.cm - 2) < 0.01 for value in [section.top_margin, section.left_margin, section.right_margin, section.bottom_margin])
    assert "Heading 1" in [p.style.name for p in word.paragraphs]
    assert all("## 第" not in p.text for p in word.paragraphs)
    assert word.styles["Normal"].font.size.pt == 11.5
    assert word.styles["Option"].paragraph_format.left_indent.cm > 0
    assert word.styles["Normal"].paragraph_format.line_spacing == 1.35
    with zipfile.ZipFile(tmp_path / "out" / "sample.docx") as archive:
        assert "word/styles.xml" in archive.namelist()
        assert b'w:type="page"' not in archive.read("word/document.xml")
    assert json.loads((tmp_path / "out" / "structured.json").read_text())["issues"]
    assert word.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "Droid Sans Fallback"


def test_long_stem_flows_across_pages(tmp_path):
    original = tmp_path / "ocr.md"
    original.write_text(FIXTURE.read_text().replace("以下较长题干", "这是用于验证跨页排版的长题干。" * 350 + "以下较长题干"))
    font = tmp_path / "font.ttf"
    font.write_bytes(fitz.Font("cjk").buffer)
    checks = generate_sample(original, tmp_path / "out", font)
    assert len(checks["pdf_pages"]) >= 3
    assert checks["content_equal"]


def test_review_edits_are_independent_and_exactly_bound_to_source():
    raw = FIXTURE.read_text()
    line = raw.splitlines().index("4. 写出几本较大的著") + 1
    edit = {"page": 6, "line": line, "original": "4. 写出几本较大的著",
            "replacement": ["4. 合成的人工核对题干。"], "verified": True,
            "reference": "synthetic-page-6.png", "reason": "synthetic fixture evidence"}
    model = structure_ocr(raw, reviewed_edits=[edit])
    question = next(b for b in model["blocks"] if b["kind"] == "question")
    assert question["parts"][0]["text"] == "合成的人工核对题干。"
    assert question["fragments"][0]["raw"] == edit["original"]
    assert model["reviewed_edits"] == [edit]
    assert model["source_sha256"] == hashlib.sha256(raw.encode()).hexdigest()
    with pytest.raises(ValueError, match="REVIEW_SOURCE_MISMATCH"):
        structure_ocr(raw, reviewed_edits=[{**edit, "original": "different"}])
    with pytest.raises(ValueError, match="INVALID_REVIEW_EVIDENCE"):
        structure_ocr(raw, reviewed_edits=[{**edit, "verified": False}])
    with pytest.raises(ValueError, match="REVIEW_SOURCE_NOT_FOUND"):
        structure_ocr(raw, reviewed_edits=[{**edit, "line": 999}])


def test_margin_removal_requires_evidence_and_keeps_original_fragment():
    raw = FIXTURE.read_text().replace("B. 直接改写不清楚的文字。", "B. 直接改写不清楚的文字。\n边栏字")
    line = raw.splitlines().index("边栏字") + 1
    evidence = {"page": 6, "line": line, "text": "边栏字", "region": "header",
                "verified": True, "reference": "synthetic-page-6.png"}
    model = structure_ocr(raw, evidence=[evidence])
    assert any(r["fragment"]["raw"] == "边栏字" and r["evidence"] == evidence for r in model["removed_headers_footers"])
    assert "边栏字" in structure_ocr(raw)["blocks"][1]["parts"][2]["text"]
    with pytest.raises(ValueError, match="INVALID_MARGIN_EVIDENCE"):
        structure_ocr(raw, evidence=[{**evidence, "reference": ""}])
    with pytest.raises(ValueError, match="MARGIN_SOURCE_MISMATCH"):
        structure_ocr(raw, evidence=[{**evidence, "text": "another"}])
