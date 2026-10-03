"""V1.1 lifecycle integration, real streamed bytes and real Python output verification."""
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace

import fitz
import openpyxl
import pytest

from conftest import register
from nas_filetools.catalog import CATEGORIES, advance_bytes, digest_file
from nas_filetools.contracts import Fault, Limits, owner
from nas_filetools.operations import diagnose, install, migrate
from nas_filetools.service import dispatch
from nas_filetools.store import Store
from nas_filetools.worker import Supervisor
from test_core import completed


def finish(store, identity, job):
    supervisor = Supervisor(store)
    try:
        for _ in range(400):
            supervisor.tick()
            result = store.status(identity, job["job_id"])
            if result["status"] not in ("QUEUED", "RUNNING"):
                return result
            time.sleep(0.05)
        raise AssertionError("Job did not complete")
    finally:
        supervisor.close()


def test_init_and_concurrent_agent_dedup_preserve_upload_events(store, identity, samples):
    marker = store.initialize("chen")/"inbox"/"keep"
    marker.write_bytes(b"keep")
    Store(store.root, store.limits)
    assert marker.read_bytes() == b"keep"
    for category in CATEGORIES:
        assert (store.root/"agents"/"chen"/"snapshots"/category).is_dir()
    def upload(i):
        return store.register(identity, f"{i:032x}", f"notes-{i}.md", samples/"notes.md", {"source_message_id": str(i)})
    with ThreadPoolExecutor(max_workers=4) as executor:
        results = list(executor.map(upload, range(1, 5)))
    assert len({r["file_id"] for r in results}) == 1
    item = store.files(identity, "list")["items"][0]
    assert len(item["uploads"]) == 4 and len({u["message_id"] for u in item["uploads"]}) == 4
    assert len(list((store.root/"agents"/"chen"/"snapshots"/"markdown").iterdir())) == 1
    (samples/"notes.md").write_text("Different same-name content")
    assert store.register(identity, "9"*32, "notes-1.md", samples/"notes.md")["file_id"] != item["file_id"]
    other = {"agent_id": "liang", "user_id": "liang", "session_hash": "b"*64}
    with pytest.raises(Fault, match="NOT_FOUND"):
        store.files(other, "select", file_id=item["file_id"])
    with pytest.raises(Fault, match="FORBIDDEN"):
        store.files({**identity, "agent_id": "stranger", "user_id": "stranger"}, "list")


def test_archive_unsupported_and_actual_type_before_parse(store, identity, samples):
    result = store.register(identity, "1"*32, "renamed.exe", samples/"text.pdf")
    assert result["category"] == "pdf" and result["extension_mismatch"]
    assert result["units"] == 0 and result["deferred_inspection"]
    assert store.inspect(identity, "1"*32)["units"] == 2
    unsupported = samples/"unknown.bin"
    unsupported.write_bytes(b"\0\xffunsupported")
    result = store.register(identity, "2"*32, unsupported.name, unsupported)
    assert result["category"] == "other" and result["kind"] == "unsupported"
    with pytest.raises(Fault, match="UNSUPPORTED_TYPE"):
        store.submit(identity, "2"*32, {"mode": "full"})
    assert store.inspect(identity)["attachments"][-1]["attachment_id"] == "2"*32


def test_ttl_read_nohit_poll_invalid_and_expired_snapshot_rebuild(store, identity, samples):
    job = completed(store, identity, samples/"notes.md")
    jid = job["job_id"]
    with store.db() as db:
        db.execute("UPDATE jobs SET expires=? WHERE id=?", (time.time()+10, jid))
    before = store.job(identity, jid)["expires"]
    store.status(identity, jid)
    store.read(identity, jid, {"keyword": "missingxyz"}, find=True)
    with pytest.raises(Fault):
        store.read(identity, jid, {"artifact_id": "bad"})
    assert store.job(identity, jid)["expires"] == before
    store.read(identity, jid, {"paragraphs": [1, 1]})
    assert store.job(identity, jid)["expires"] > before+100
    with store.db() as db:
        db.execute("UPDATE jobs SET expires=0 WHERE id=?", (jid,))
    with pytest.raises(Fault, match="CACHE_EXPIRED"):
        store.read(identity, jid, {})
    store.cleanup()
    assert store.inspect(identity, "1"*32)["sha256"]
    rebuilt = store.submit(identity, "1"*32, {"mode": "full"})
    assert rebuilt["job_id"] != jid
    new_session = {**identity, "session_hash": "c"*64}
    assert store.inspect(new_session)["attachments"] == []
    file_id = store.files(new_session, "list")["items"][0]["file_id"]
    selected = store.files(new_session, "select", file_id=file_id)
    assert selected["attachment_id"] != "1"*32
    with pytest.raises(Fault, match="NOT_FOUND"):
        store.status(new_session, rebuilt["job_id"])


def test_saved_bundle_links_versions_failure_cleanup_and_delete(store, identity, samples, monkeypatch):
    job = completed(store, identity, samples/"ordered.docx")
    assert store.files(identity, "offer_save", job_id=job["job_id"])["ask_once"]
    assert not store.files(identity, "offer_save", job_id=job["job_id"])["ask_once"]
    import nas_filetools.workspace as module
    real = module.shutil.copyfile
    monkeypatch.setattr(module.shutil, "copyfile", lambda *_: (_ for _ in ()).throw(OSError("full")))
    with pytest.raises(OSError):
        store.save_result(identity, job["job_id"])
    assert (store.task(job["job_id"])/"content.md").is_file()
    monkeypatch.setattr(module.shutil, "copyfile", real)
    saved = store.save_result(identity, job["job_id"])
    another = store.save_result(identity, job["job_id"])
    assert saved["saved_id"] != another["saved_id"]
    with store.db() as db:
        folder = store.root/db.execute("SELECT path FROM saved WHERE id=?", (saved["saved_id"],)).fetchone()[0]
        db.execute("UPDATE jobs SET expires=0 WHERE id=?", (job["job_id"],))
    md = (folder/"content.md").read_text(encoding="utf-8")
    assert "artifact:" not in md and "assets/" in md
    assert list((folder/"assets").glob("*.png")) and (folder/"sources.json").exists()
    store.cleanup()
    assert "Before the table" in store.saved_read(identity, saved["saved_id"], 0)["text"]
    original = store.files(identity, "list")["items"][0]
    store.files(identity, "remove_reference", attachment_id="1"*32)
    assert store.files(identity, "list")["items"][0]["file_id"] == original["file_id"]
    store.files(identity, "delete_original", file_id=original["file_id"])
    assert store.saved_read(identity, saved["saved_id"], 0)["original_state"] == "DELETED"
    assert store.files(identity, "list")["items"] == []
    assert (samples/"ordered.docx").is_file() and not (store.root/"index").exists()


def test_active_jobs_protect_original_and_cache(store, identity, samples):
    attachment = register(store, identity, samples/"notes.md")
    job = store.submit(identity, "1"*32, {"mode": "full"})
    with store.db() as db:
        db.execute("UPDATE jobs SET expires=0")
    assert store.cleanup()["removed_jobs"] == []
    for action, params in [("delete_original", {"file_id": attachment["file_id"]}),
                           ("remove_reference", {"attachment_id": "1"*32}), ("delete_cache", {"job_id": job["job_id"]})]:
        with pytest.raises(Fault, match="FILE_BUSY"):
            store.files(identity, action, **params)


def test_11000_chinese_minutes_fit_central_control_limit(store, identity, samples):
    job = completed(store, identity, samples/"notes.md")
    params = {"job_id": job["job_id"], "text": "中"*11000, "source_segments": [1]}
    payload = json.dumps({"identity": identity, "operation": "save_minutes", "params": params}, ensure_ascii=False).encode()
    assert 32768 < len(payload) < store.limits.max_control_bytes
    result = dispatch(store, identity, "save_minutes", params)
    read = store.read(identity, job["job_id"], {"artifact_id": result["artifact_id"]})
    assert len(read["text"]) == 11000


def test_start_failure_terminal_and_queue_ttl_not_processing_timeout(store, identity, samples, monkeypatch):
    import nas_filetools.worker as worker
    register(store, identity, samples/"notes.md")
    job = store.submit(identity, "1"*32, {"mode": "full"})
    with store.db() as db:
        db.execute("UPDATE jobs SET expires=0")
    class Broken:
        def start(self):
            raise OSError("spawn")
    monkeypatch.setattr(worker.multiprocessing, "get_context", lambda _: type("Context", (), {"Process": lambda **_: Broken()})())
    supervisor = Supervisor(store)
    supervisor.tick()
    assert store.status(identity, job["job_id"])["status"] == "FAILED"
    assert not supervisor.processes


def test_scheduler_fatal_error_stops_service_instead_of_healthy_dead_thread():
    import threading
    from nas_filetools.service import scheduler_loop
    stop, state = threading.Event(), {"error": None, "last_tick": 0}
    class Broken:
        def tick(self):
            raise RuntimeError("database failure")
    thread = threading.Thread(target=scheduler_loop, args=(Broken(), stop, state))
    thread.start()
    thread.join(2)
    assert stop.is_set() and not thread.is_alive() and state["error"] == "RuntimeError"


def test_real_201_mib_inbox_stream_and_4gib_counter_only(store, identity):
    inbox = store.initialize("chen")/"inbox"
    file = inbox/"large.bin"
    # Write every byte: neither sparse nor mocked. SHA/copy use 1 MiB buffers.
    block = bytes(range(256))*4096
    with file.open("wb") as output:
        for _ in range(201):
            output.write(block)
    assert file.stat().st_size == 201*1024**2
    item = store.inbox_list(identity)["items"][0]
    with pytest.raises(Fault, match="INBOX_INCOMPLETE"):
        store.inbox_register(identity, item["inbox_id"])
    for incomplete in ("{", "[]", '{"bytes":true,"sha256":"incomplete"}'):
        (inbox/"large.bin.ready").write_text(incomplete)
        with pytest.raises(Fault, match="INBOX_INCOMPLETE"):
            store.inbox_register(identity, item["inbox_id"])
    original_limits = store.limits
    store.limits = replace(store.limits, max_receive_bytes=200*1024**2, max_process_bytes=200*1024**2)
    with pytest.raises(Fault, match="RECEIVE_SIZE_LIMIT"):
        store.inbox_register(identity, item["inbox_id"])
    store.limits = original_limits
    (inbox/"large.bin.ready").write_text(json.dumps({"bytes": file.stat().st_size, "sha256": digest_file(file)}))
    result = store.inbox_register(identity, item["inbox_id"])
    assert result["bytes"] == file.stat().st_size and result["sha256"] == digest_file(file)
    assert file.is_file()
    assert advance_bytes(2**32-1024, 1024, 4*1024**3) == 2**32
    with pytest.raises(Fault, match="RECEIVE_SIZE_LIMIT"):
        advance_bytes(2**32, 1, 4*1024**3)


def test_real_python_xlsx_verified_copy_and_original_unchanged(tmp_path, identity):
    store = Store(tmp_path/"state", replace(Limits(), script_isolation="development"))
    file = tmp_path/"budget.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Amount"])
    sheet.append([12])
    sheet.append([5])
    book.create_sheet("Notes")["A1"] = "Preserve"
    book.save(file)
    before = digest_file(file)
    store.register(identity, "1"*32, file.name, file)
    script = """import openpyxl
book = openpyxl.load_workbook('input.xlsx')
sheet = book['Budget']
sheet['B1'] = 'Double'
for row in range(2, sheet.max_row+1):
    sheet.cell(row, 2, sheet.cell(row, 1).value*2)
book.save('result.xlsx')
"""
    params = {"attachment_id": "1"*32, "description": "Add Double numeric column; preserve Notes", "code": script,
              "outputs": ["result.xlsx"], "checks": [{"file": "result.xlsx", "sheet": "Budget", "cell": "B2", "equals": 24,
              "min_rows": 3, "min_columns": 2}, {"file": "result.xlsx", "sheet": "Notes", "cell": "A1", "equals": "Preserve"}]}
    job = dispatch(store, identity, "python", params)
    result = finish(store, identity, job)
    assert result["status"] == "SUCCEEDED", result
    output = store.task(job["job_id"])/next(a["file"] for a in result["artifacts"] if a["kind"] == "output")
    modified = openpyxl.load_workbook(output)
    assert modified["Budget"]["B3"].value == 10
    assert digest_file(file) == before
    code_artifact = next(a for a in result["artifacts"] if a["kind"] == "script")
    assert "load_workbook" in store.read(identity, job["job_id"], {"artifact_id": code_artifact["artifact_id"]})["text"]
    saved = store.save_result(identity, job["job_id"], next(a["artifact_id"] for a in result["artifacts"] if a["kind"] == "output"))
    store.files(identity, "delete_cache", job_id=job["job_id"])
    assert store.saved_read(identity, saved["saved_id"], 0)["artifact"]["file"].endswith("/result.xlsx")
    assert any("DEVELOPMENT MODE" in w for w in result["warnings"])


@pytest.mark.parametrize("code,expected", [("raise RuntimeError('failure')", "SCRIPT_FAILED"),
    ("pass", "OUTPUT_NOT_VERIFIED"), ("while True: pass", "SCRIPT_TIMEOUT")])
def test_script_failure_missing_output_and_timeout(tmp_path, identity, code, expected):
    store = Store(tmp_path/"state", replace(Limits(), script_isolation="development", script_timeout_seconds=1))
    file = tmp_path/"notes.txt"
    file.write_text("input")
    store.register(identity, "1"*32, file.name, file)
    job = dispatch(store, identity, "python", {"attachment_id": "1"*32, "code": code, "description": "failure case", "outputs": ["result.txt"]})
    result = finish(store, identity, job)
    assert result["status"] == "FAILED" and result["failures"][0]["code"] == expected


def test_missing_script_boundary_is_explicit_on_windows(store, identity):
    if sys.platform != "win32":
        pytest.skip("Windows capability boundary")
    with pytest.raises(Fault, match="SCRIPT_ISOLATION_UNAVAILABLE"):
        dispatch(store, identity, "python", {"attachment_id": "1"*32, "code": "pass", "description": "test", "outputs": ["result.txt"]})


def test_real_script_output_quota_failure_preserves_original(tmp_path, identity):
    store = Store(tmp_path/"state", replace(Limits(), script_isolation="development", script_output_bytes=1024))
    original = tmp_path/"notes.txt"
    original.write_text("immutable source")
    original_sha = digest_file(original)
    store.register(identity, "1"*32, original.name, original)
    job = dispatch(store, identity, "python", {"attachment_id": "1"*32,
        "code": "from pathlib import Path\nPath('result.txt').write_bytes(b'x'*(3*1024**2))",
        "description": "Verify bounded output failure", "outputs": ["result.txt"]})
    result = finish(store, identity, job)
    assert result["status"] == "FAILED" and result["failures"][0]["code"] == "SCRIPT_OUTPUT_LIMIT"
    assert not result["artifacts"] and digest_file(original) == original_sha


def test_large_pdf_last_page_and_model_config_fingerprint(store, identity, tmp_path):
    file = tmp_path/"large.pdf"
    with fitz.open() as pdf:
        for index in range(601):
            pdf.new_page().insert_text((30, 50), f"Cedar manual page {index+1}. Digital text with enough length.")
        pdf.save(file)
    store.register(identity, "1"*32, file.name, file)
    inspected = store.inspect(identity, "1"*32)
    assert inspected["units"] == 601 and len(inspected["preview"]) == 3
    job = store.submit(identity, "1"*32, {"mode": "range", "pages": [601, 601]})
    result = finish(store, identity, job)
    assert result["status"] == "SUCCEEDED"
    assert store.read(identity, job["job_id"], {})["evidence"][0]["source"]["page"] == 601
    different = store.submit(identity, "1"*32, {"mode": "range", "pages": [600, 600]})
    assert different["job_id"] != job["job_id"]


def test_local_install_diagnosis_twice_and_v1_migration(tmp_path, identity):
    root = tmp_path/"new"
    configuration = tmp_path/"install.json"
    configuration.write_text(json.dumps({"data_root": str(root), "agents": ["chen"], "bindings": []}))
    assert install(configuration, local_only=True)["status"] == "INSTALL_PLAN"
    assert not root.exists()
    assert install(configuration, apply=True, local_only=True)["status"] == "LOCAL_WORKSPACE_READY"
    keep = root/"agents"/"chen"/"inbox"/"keep"
    keep.write_text("keep")
    install(configuration, apply=True, local_only=True)
    assert keep.read_text() == "keep"
    assert diagnose(configuration, local_only=True)["checks"]["runtime"] == "NOT_CHECKED_LOCAL_ONLY"
    old = tmp_path/"old"
    old.mkdir()
    scope = owner(identity)
    path = old/"originals"/scope/("1"*32+".txt")
    path.parent.mkdir(parents=True)
    path.write_text("old content")
    import sqlite3
    with sqlite3.connect(old/"jobs.sqlite3") as db:
        db.execute("CREATE TABLE attachments(id,owner,filename,sha,size,info,created)")
        db.execute("INSERT INTO attachments VALUES(?,?,?,?,?,?,?)", ("1"*32, scope, "old.txt", digest_file(path), path.stat().st_size,
                   json.dumps({"kind": "text", "units": 1, "warnings": []}), 0))
    mapping = tmp_path/"scopes.json"
    mapping.write_text(json.dumps({scope: identity}))
    assert migrate(old, root, Limits(), mapping)["status"] == "MIGRATION_PLAN"
    result = migrate(old, root, Limits(), mapping, apply=True)
    assert result["status"] == "MIGRATED" and path.read_text() == "old content"
    assert Store(root, Limits()).inspect(identity, "1"*32)["filename"] == "old.txt"
    assert list(root.glob("v1-metadata-*.sqlite3"))
    assert Store(root, Limits()).files(identity, "list")["items"][0]["first_received_at"].startswith("1970-01-01")


def test_queue_rechecks_disk_and_oversized_engine_log_is_terminal(store, identity, samples):
    register(store, identity, samples/"notes.md")
    job = store.submit(identity, "1"*32, {"mode": "full"})
    # Disk availability may change after submission. Check the actual queued-start guard.
    store.limits = replace(store.limits, cache_quota_bytes=64*1024**2)
    supervisor = Supervisor(store)
    supervisor.tick()
    result = store.status(identity, job["job_id"])
    assert result["status"] == "FAILED" and result["failures"][0]["code"] == "DISK_QUOTA"
    assert not supervisor.processes
    store.limits = replace(store.limits, cache_quota_bytes=20*1024**3, max_engine_log_bytes=128)
    second = store.submit(identity, "1"*32, {"mode": "full"})
    (store.task(second["job_id"])/"engine.log").write_bytes(b"x"*129)
    with store.db() as db:
        db.execute("UPDATE jobs SET status='RUNNING' WHERE id=?", (second["job_id"],))
    class ExitedProcess:
        def is_alive(self):
            return False
        def join(self, *_args):
            pass
    supervisor.processes[second["job_id"]] = (ExitedProcess(), time.monotonic())
    supervisor.tick()
    result = store.status(identity, second["job_id"])
    assert result["status"] == "FAILED" and result["failures"][0]["code"] == "ENGINE_LOG_LIMIT"
    supervisor.close()


def test_interrupted_exact_delete_recovers_and_preserves_saved(store, identity, samples):
    job = completed(store, identity, samples/"notes.md")
    file_id = store.inspect(identity, "1"*32)["file_id"]
    saved = store.save_result(identity, job["job_id"], None)
    # Simulate a committed deletion intent followed by a process stop before unlink.
    with store.db() as db:
        db.execute("UPDATE files SET state='DELETING' WHERE id=?", (file_id,))
    restarted = Store(store.root, store.limits)
    assert restarted.files(identity, "list")["items"] == []
    assert restarted.saved_read(identity, saved["saved_id"], 0)["original_state"] == "DELETED"
    assert (samples/"notes.md").exists()
    assert not list((store.root/"agents"/"chen"/"snapshots"/"markdown").iterdir())
