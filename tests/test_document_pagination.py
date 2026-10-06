"""Synthetic boundaries, with actual Story and optional native Word rendering."""
import json
import os
import shutil
import subprocess

import fitz
import pytest
from docx import Document
from docx.oxml.ns import qn

from nas_filetools.document_layout_checks import check_layout
from nas_filetools.document_sample import StoryLayout, validate_pair, write_docx, write_pdf


def question(stem="短题完整放置验证。", option="一个完整的短选项。", last=None):
    return [("Question", "207.\u00a0" + stem), ("Option", "A. " + option),
            ("Option", "B. " + option), ("Option", "C. " + option),
            ("Option", "D. " + (last or option)), ("Source", "原PDF：第88页")]


@pytest.fixture
def font(tmp_path):
    path = tmp_path / "font.ttf"
    path.write_bytes(fitz.Font("cjk").buffer)
    return path


def bottom_start_items(font):
    layout = StoryLayout(font)
    group = question()
    title = [("Title", "合成分页边界")]
    filler = "这是已存在的前题内容，用于制造自然的页底剩余空间。"
    # Determine the boundary using actual layout; no magic page or question IDs.
    count = 1
    while layout.height(title + [("Unresolved", filler * count)]) < layout.where.height - layout.height(group) / 2:
        count += 1
    return title + [("Unresolved", filler * count), ("Heading 2", "第二章 合成标题边界")] + group


def write_pair(directory, items, font):
    directory.mkdir(exist_ok=True)
    docx, pdf = directory / "sample.docx", directory / "sample.pdf"
    write_docx(docx, items, font)
    write_pdf(pdf, items, font)
    validate_pair(docx, pdf, items)
    return docx, pdf


@pytest.mark.parametrize("case", ["bottom-short", "long-stem", "long-option", "long-final-option"])
def test_actual_pdf_generic_pagination(case, font, tmp_path):
    if case == "bottom-short":
        items = bottom_start_items(font)
    elif case == "long-stem":
        items = [("Title", "合成分页边界")] + question(stem="这是允许自然跨页的超长题干，不应全题绑定。" * 230)
    elif case == "long-option":
        items = [("Title", "合成长选项")] + question(option="长选项的完整文字应保留，允许合理分成多个页面。" * 140)
    else:
        items = [("Title", "合成末选项")] + question(last="末选项很长，来源仍须跟随最后两行，而不是单独占页。" * 200)
    docx, pdf = write_pair(tmp_path, items, font)
    evidence = check_layout(pdf, items, font)
    assert evidence["passed"], evidence["failures"]
    q = evidence["questions"][0]
    assert q["source_follows_last_option"]
    if case == "bottom-short":
        assert q["pages"] == [2]
        assert evidence["pages"][0]["bottom_blank_pt"] < StoryLayout(font).short_limit
    else:
        assert len(q["pages"]) > 1
        if case in ("long-option", "long-final-option"):
            split = [o for o in q["options"] if len(o["pages"]) > 1]
            assert split
            assert all(min(o["lines_per_page"].values()) >= 2 for o in split)
    assert len(Document(docx).paragraphs) == len(items)


@pytest.mark.skipif(shutil.which("libreoffice") is None, reason="independent Word renderer not installed")
@pytest.mark.parametrize("case", ["bottom-short", "long-stem", "long-option", "long-final-option"])
def test_actual_word_native_pagination(case, font, tmp_path):
    items = bottom_start_items(font) if case == "bottom-short" else [("Title", "合成Word分页边界")] + question(
        stem="原生孤行控制允许长题跨页，保留完整文字。" * 230 if case == "long-stem" else "短题干。",
        option="长选项完整保留文字，允许按原生规则跨页。" * 150 if case == "long-option" else "完整短选项。",
        last="长末选项允许跨页，来源保持在末尾同页。" * 230 if case == "long-final-option" else None)
    docx, pdf = write_pair(tmp_path, items, font)
    exported = tmp_path / "word"
    exported.mkdir()
    environment = {**os.environ, "XDG_CACHE_HOME": str(tmp_path / "cache"), "GSETTINGS_BACKEND": "memory"}
    subprocess.run(["libreoffice", "-env:UserInstallation=" + (tmp_path / "profile").as_uri(),
                    "--headless", "--convert-to", "pdf", "--outdir", str(exported), str(docx)],
                   check=True, timeout=120, env=environment, capture_output=True)
    validate_pair(docx, exported / "sample.pdf", items)
    evidence = check_layout(exported / "sample.pdf", items, font)
    assert evidence["passed"], evidence["failures"]
    q = evidence["questions"][0]
    assert q["source_follows_last_option"]
    assert len(q["pages"]) == 1 if case == "bottom-short" else len(q["pages"]) > 1
    if case == "long-stem":
        assert q["stem_opening_lines"] >= 2
    if case == "long-option":
        assert q["options"][0]["pages"][0] == q["pages"][0]
        assert all(min(o["lines_per_page"].values()) >= 2 for o in q["options"])
    if case == "long-final-option":
        last = q["options"][-1]
        assert len(last["pages"]) > 1
        assert min(last["lines_per_page"].values()) >= 2
    (tmp_path / "word-evidence.json").write_text(json.dumps(evidence))


def test_black_hierarchy_footer_and_bounded_word_keeps(font, tmp_path):
    items = [("Title", "合成标题"), ("Heading 1", "第一部分"), ("Subject", "合成学科"),
             ("Heading 3", "一、单项选择题"), ("Heading 2", "导论")] + question() + [
                 ("Question", "208.\u00a0" + "长题干允许跨页。" * 240),
                 ("Option", "A. 短选项"), ("Option", "B. 短选项"), ("Option", "C. 短选项"),
                 ("Option", "D. 短选项"), ("Source", "原PDF：第89页")]
    docx, pdf = write_pair(tmp_path, items, font)
    word = Document(docx)
    assert [p.style.name for p in word.paragraphs[:5]] == [s for s, text in items[:5]]
    for style in ("Title", "Heading 1", "Subject", "Heading 2", "Heading 3"):
        color = word.styles[style].element.rPr.find(qn("w:color"))
        assert color.get(qn("w:val")) == "000000" and color.get(qn("w:themeColor")) is None
    short_last_source = word.paragraphs[10]
    assert short_last_source.style.name == "Source" and not short_last_source.paragraph_format.keep_with_next
    long_stem = word.paragraphs[11]
    assert not long_stem.paragraph_format.keep_with_next and not long_stem.paragraph_format.keep_together
    assert word.styles["Normal"].paragraph_format.widow_control
    footer = word.sections[0].footer.paragraphs[0]
    assert footer._p.find(qn("w:fldSimple")).get(qn("w:instr")) == "PAGE"
    assert check_layout(pdf, items, font)["passed"]


def test_previous_content_comparison_excludes_only_declared_display_changes(font, tmp_path):
    from nas_filetools.document_layout_checks import compare_docx_content
    before = [("Title", "合成标题"), ("Heading 1", "第一章")] + question()
    before[-1] = ("Source", "来源：原PDF第88页；OCR行1–8")
    after = [("Title", "合成标题"), ("Heading 2", "第一章")] + question()
    old, new = tmp_path / "old.docx", tmp_path / "new.docx"
    write_docx(old, before, font)
    write_docx(new, after, font)
    assert compare_docx_content(old, new)["exact_question_option_text_and_order"]
    after[-2] = ("Option", "D. 不允许默默改变文字。")
    write_docx(new, after, font)
    with pytest.raises(ValueError, match="PREVIOUS_DOCUMENT_CONTENT_CHANGED"):
        compare_docx_content(old, new)
