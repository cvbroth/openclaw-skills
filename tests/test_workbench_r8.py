"""Synthetic Word packages and real office pagination, no business text."""

import copy
import json
import zipfile
import pytest
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from nas_filetools.standalone.word_import import extract
from nas_filetools.standalone.word_styles import validate_styles, apply_styles
from nas_filetools.standalone.template_library import TemplateLibrary
from nas_filetools.standalone.source_pagination import options, group_pages
from nas_filetools.standalone.documents import generate


def sample(path, multi=False):
    d = Document()
    d.styles["Normal"].font.name = "Missing-Example-Font"
    d.styles["Normal"].font.size = Pt(12)
    d.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hint"), "eastAsia")
    d.styles["Normal"].paragraph_format.line_spacing = Pt(18)
    base = d.styles.add_style("SyntheticBase", 1)
    base.font.size = Pt(17)
    d.styles["Heading 1"].base_style = base
    d.styles["Heading 1"].font.size = None
    d.add_paragraph("明确的合成文字，不是原试题册。")
    cols = d.sections[0]._sectPr.find(qn("w:cols"))
    cols.set(qn("w:num"), "2")
    if multi:
        d.add_section(WD_SECTION_START.NEW_PAGE).left_margin = Cm(3)
    d.save(path)
    return path


def test_inheritance_exact_multiple_and_dotx(tmp_path):
    p = sample(tmp_path / "sample.docx")
    r = extract(p, p.name)
    assert r["paragraphs"]["Normal"]["line_mode"] == "exact"
    assert r["paragraphs"]["Normal"]["line_value"] == 18
    assert r["paragraphs"]["Heading 1"]["size_pt"] == 17
    assert any("分栏" in v for v in r["unsupported"])
    dotx = tmp_path / "sample.dotx"
    with zipfile.ZipFile(p) as src, zipfile.ZipFile(dotx, "w") as dst:
        for n in src.namelist():
            raw = src.read(n)
            if n == "[Content_Types].xml":
                raw = raw.replace(
                    b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml"
                )
            dst.writestr(n, raw)
    assert extract(dotx, dotx.name)["paragraphs"] == r["paragraphs"]
    d = Document(p)
    d.styles["Normal"].paragraph_format.line_spacing = 1.25
    d.save(p)
    assert extract(p, p.name)["paragraphs"]["Normal"]["line_value"] == 1.25


def test_multisection_font_selection_snapshot(tmp_path):
    p = sample(tmp_path / "multi.docx", True)
    lib = TemplateLibrary(tmp_path / "templates")
    data = lib.import_word(p, p.name)
    assert data["document"] is None and data["report"]["section_selection_required"]
    assert "Missing-Example-Font" in data["report"]["missing_fonts"]
    assert "eastAsia" not in data["report"]["missing_fonts"]
    d = lib.select_word_import(data["report"]["import_id"], 1, "Droid Sans Fallback")["document"]
    assert d["word_styles"]["section"]["left_cm"] == pytest.approx(3, abs=0.001)
    lib.save(d)
    actual = lib.load(d["id"], d["version"], True)
    assert actual["word_styles"] == d["word_styles"]
    original = lib.root / "imports" / data["report"]["import_id"] / "original.docx"
    assert original.read_bytes() == p.read_bytes()
    broken = copy.deepcopy(d)
    broken["version"] = "1.0.1"
    broken["import_source"]["font_selection_confirmed"] = False
    with pytest.raises(ValueError, match="font_selection"):
        lib.save(broken)


def test_unsafe_packages_and_field_validation(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_bytes(b"not zip")
    with pytest.raises(ValueError, match="无效"):
        extract(p, p.name)
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("../escape", "invalid")
    with pytest.raises(ValueError, match="路径"):
        extract(p, p.name)
    with pytest.raises(ValueError, match="paragraphs"):
        validate_styles({"section": {}, "paragraphs": []})
    with pytest.raises(ValueError, match="行距"):
        validate_styles({"section": {}, "paragraphs": {"Normal": {"line_mode": "exact"}}})
    with pytest.raises(ValueError, match="页边距"):
        validate_styles({"section": {"width_cm": 10, "left_cm": 6, "right_cm": 6}, "paragraphs": {}})
    with pytest.raises(ValueError):
        options({"minimum_font_pt": 2})
    with pytest.raises(ValueError, match="MAPPING"):
        group_pages([{"items": [], "physical_page": 1}])


def test_actual_word_styles_execution(tmp_path):
    p = sample(tmp_path / "sample.docx")
    extension = {
        "section": {"width_cm": 29.7, "height_cm": 21, "left_cm": 2.1},
        "paragraphs": {
            "Normal": {"line_mode": "exact", "line_value": 19, "size_pt": 13, "first_line_cm": 0.6}
        },
    }
    apply_styles(p, extension)
    d = Document(p)
    assert d.sections[0].page_width.cm == pytest.approx(29.7, abs=0.001)
    assert d.styles["Normal"].paragraph_format.line_spacing.pt == 19
    assert d.styles["Normal"].font.size.pt == 13


def request(tmp_path, text=None, continuation=False):
    t = TemplateLibrary(tmp_path / "library").load("questions-zh-cn", "1.0.0")
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    return {
        "output": str(out),
        "title": "禁止插入正文的项目大标题",
        "template": {"id": "questions-zh-cn", "version": "1.0.0"},
        "template_data": t,
        "pagination": {"mode": "source-pages", "allow_continuation": continuation},
        "pages": [
            {
                "physical_page": 1,
                "source_pdf_physical_page": 23,
                "text": text
                if text is not None
                else "## 合成标题\n\n完整正文。\n\nA. 保留选项。\n\nB. 独立审计。",
            },
            {"physical_page": 2, "source_pdf_physical_page": 24, "text": ""},
        ],
    }


def test_real_source_pages_and_blank(tmp_path):
    r = request(tmp_path)
    r["pagination"]["font_fit"] = True
    generate(r)
    report = json.loads((tmp_path / "out/pagination-report.json").read_text())
    assert report["verified"]
    assert [g["output_pages"] for g in report["groups"]] == [[1], [2]]
    d = Document(tmp_path / "out/document.docx")
    assert r["title"] not in "\n".join(p.text for p in d.paragraphs)
    assert json.loads((tmp_path / "out/layout-validation.json").read_text())["content_equal"]


def test_real_overflow_conflict_and_explicit_continuation(tmp_path):
    r = request(
        tmp_path, "\n\n".join("合成超长材料" + str(i) + "。" + "长段落完整保留。" * 15 for i in range(70))
    )
    with pytest.raises(ValueError, match="SOURCE_PAGE_OVERFLOW"):
        generate(r)
    report = json.loads((tmp_path / "out/pagination-report.json").read_text())
    assert report["groups"][0]["attempts"][-1]["actual_output_pages"] > 1
    r["pagination"]["allow_continuation"] = True
    generate(r)
    report = json.loads((tmp_path / "out/pagination-report.json").read_text())
    assert report["verified"] and len(report["groups"][0]["output_pages"]) > 1


def test_real_question_tables_keep_text_and_mapping(tmp_path):
    r = request(tmp_path, "1. 合成完整题干，明确四个选项。\n\nA. 甲。\nB. 乙。\nC. 丙。\nD. 丁。")
    generate(r)
    report = json.loads((tmp_path / "out/pagination-report.json").read_text())
    assert report["verified"] and report["groups"][0]["output_pages"] == [1]
    checks = json.loads((tmp_path / "out/layout-validation.json").read_text())
    assert checks["content_equal"]


def test_group_order_and_repeated_same_source():
    a = {"source_pdf_physical_page": 23, "items": [("Normal", "甲")]}
    groups = group_pages([a, {**a, "items": [("Normal", "乙")]}])
    assert len(groups) == 1 and len(groups[0]["items"]) == 2
    with pytest.raises(ValueError, match="ORDER_CONFLICT"):
        group_pages([a, {**a, "source_pdf_physical_page": 22}])


def test_macro_and_entity_rejected(tmp_path):
    p = sample(tmp_path / "sample.docx")
    bad = tmp_path / "macro.docx"
    with zipfile.ZipFile(p) as src, zipfile.ZipFile(bad, "w") as dst:
        for n in src.namelist():
            dst.writestr(n, src.read(n))
        dst.writestr("word/vbaProject.bin", b"not executed")
    with pytest.raises(ValueError, match="宏"):
        extract(bad, bad.name)
    from nas_filetools.standalone.word_import import xml

    with zipfile.ZipFile(tmp_path / "entity.zip", "w") as dst:
        dst.writestr(
            "x.xml", '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///nonexistent">]><x>&e;</x>'.encode("utf-16")
        )
    with zipfile.ZipFile(tmp_path / "entity.zip") as src, pytest.raises(ValueError, match="DTD"):
        xml(src, "x.xml")
