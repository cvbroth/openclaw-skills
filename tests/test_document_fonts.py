"""Actual fallback coverage and portable Word font payload, never text deletion."""

import hashlib
import uuid
from zipfile import ZipFile
from lxml import etree
import fitz
import pytest
from docx import Document
from nas_filetools.document_fonts import fallback_records, font_runs, font_audit
from nas_filetools.document_entry import FONT_PATH
from nas_filetools.document_sample import write_docx, document_paragraphs, option_layout_items
from nas_filetools.document_templates import load_template
from nas_filetools.standalone.word_styles import apply_styles
from nas_filetools.standalone.preview_worker import export_docx


def test_real_missing_glyphs_use_pinned_fonts_without_changing_text(tmp_path):
    text = "原文⸺保留📖完整"
    runs = font_runs(text, FONT_PATH)
    assert "".join(t for _, t in runs) == text
    assert ("Noto Sans", "⸺") in runs
    assert ("Noto Emoji", "📖") in runs
    audit = font_audit(
        [
            {
                "source_pdf_physical_page": 29,
                "source_locator": {"kind": "block", "id": "stable"},
                "items": [("Normal", text)],
            }
        ],
        FONT_PATH,
    )
    assert all(
        x["source_page"] == 29 and x["source_locator"]["id"] == "stable" for x in audit["substitutions"]
    )
    path = tmp_path / "symbols.docx"
    write_docx(path, [("Normal", text)], FONT_PATH)
    apply_styles(path, None, materialize=True)
    doc = Document(path)
    assert doc.paragraphs[0].text == text
    assert [r.font.name for r in doc.paragraphs[0].runs if r.text == "📖"] == ["Noto Emoji"]
    with ZipFile(path) as z:
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        table = etree.fromstring(z.read("word/fontTable.xml"))
        relations = etree.fromstring(z.read("word/_rels/fontTable.xml.rels"))
        for entry in fallback_records():
            node = next(e for e in table if e.get(w + "name") == entry["name"])
            embed = node.find(w + "embedRegular")
            rid = embed.get(rel + "id")
            target = next(e.get("Target") for e in relations if e.get("Id") == rid)
            raw = bytearray(z.read("word/" + target))
            key = uuid.UUID(embed.get(w + "fontKey")).bytes[::-1]
            for i in range(32):
                raw[i] ^= key[i % 16]
            assert hashlib.sha256(raw).hexdigest() == entry["sha256"]
    pdf = export_docx(path, tmp_path)
    with fitz.open(pdf) as rendered:
        assert text in rendered[0].get_text().replace("\n", "")
        names = [f[3] for f in rendered[0].get_fonts()]
        assert any("NotoEmoji" in n for n in names) and any("NotoSans" in n for n in names)
        assert all(
            glyph != 0
            for span in rendered[0].get_texttrace()
            for code, glyph, _, _ in span["chars"]
            if not chr(code).isspace()
        )


def test_missing_outside_chain_explains_character_and_remedy():
    with pytest.raises(ValueError, match=r"U\+0378.*安装"):
        font_runs("\u0378", FONT_PATH)


def test_option_pairing_preserves_text_and_long_options(tmp_path):
    from nas_filetools.standalone.content_structure import RichText

    text = RichText(
        [{"text": "题干\nA. 短选项甲\nB. 短选项乙\nC. " + "很长的选项" * 40 + "\nD. 最后一项", "bold": True}]
    )
    items = [("Normal", text)]
    split = option_layout_items(items)
    assert "".join(str(t) for _, t in split).replace("\n", "") == text.replace("\n", "")
    assert all(t.runs[0]["bold"] for _, t in split)
    t = load_template()["parameters"]
    t["option_columns"] = 2
    path = tmp_path / "options.docx"
    write_docx(path, items, FONT_PATH, t)
    d = Document(path)
    assert "".join(p.text for p in document_paragraphs(d)).replace("\n", "") == text.replace("\n", "")
    assert len(d.tables) == 1
    assert len(d.tables[0].columns) == 2
    assert any("很长的选项" * 40 in p.text for p in d.paragraphs)


def test_long_first_stem_is_not_used_as_word_title(tmp_path):
    text = "长题干保持完整。" * 80
    path = tmp_path / "long-stem.docx"
    write_docx(path, [("Normal", text)], FONT_PATH)
    d = Document(path)
    assert d.paragraphs[0].text == text
    assert d.core_properties.title == "FileTools 导出文档"


def test_fullwidth_option_indentation_is_preserved_and_can_pair(tmp_path):
    text = "题干\n　A．甲\n　B．乙\n　C．丙\n　D．丁"
    t = load_template()["parameters"]
    t["option_columns"] = 2
    path = tmp_path / "fullwidth.docx"
    write_docx(path, [("Normal", text)], FONT_PATH, t)
    d = Document(path)
    assert len(d.tables) == 2
    assert "".join(p.text for p in document_paragraphs(d)) == text.replace("\n", "")


def test_literal_markdown_keeps_markers_but_removes_blank_line_layout_cost(tmp_path):
    import re

    text = "# 保留原文字面标题\n\n**103.** 长题干\n\nA. 甲\n\nB. 乙\n\nC. 丙\n\nD. 丁\n\n**104.** 后题干\n\nA. 子\n\nB. 丑\n\nC. 寅\n\nD. 卯"
    t = load_template()["parameters"]
    t["option_columns"] = 2
    path = tmp_path / "literal.docx"
    write_docx(path, [("Normal", text)], FONT_PATH, t)
    d = Document(path)
    actual = "".join(p.text for p in document_paragraphs(d))
    assert re.sub(r"\s", "", actual) == re.sub(r"\s", "", text)
    assert actual.startswith("# 保留原文字面标题**103.**")
    assert len(d.tables) == 4


def test_bullet_and_fullwidth_option_labels_pair_without_deleting_markers(tmp_path):
    text = "题干\n- Ａ．甲\n- Ｂ．乙\n- Ｃ．丙\n- Ｄ．丁"
    t = load_template()["parameters"]
    t["option_columns"] = 2
    path = tmp_path / "bullets.docx"
    write_docx(path, [("Normal", text)], FONT_PATH, t)
    d = Document(path)
    assert len(d.tables) == 2
    assert "".join(p.text for p in document_paragraphs(d)) == text.replace("\n", "")
