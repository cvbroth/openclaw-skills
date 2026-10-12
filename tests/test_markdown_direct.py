"""Actual office export of reader Markdown; strict mapping, overflow and revisions."""

import copy
import json
import uuid
from unittest.mock import patch
import fitz
import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from nas_filetools.artifact_project import digest
from nas_filetools.standalone.document_exports import export_options
from nas_filetools.standalone.documents import generate
from nas_filetools.standalone.markdown_direct import direct_template
from test_template_exports import project as template_project, terminal


@pytest.fixture(name="project")
def mapped_project_fixture(tmp_path):
    yield from template_project.__wrapped__(tmp_path)


def revise(m, pid, aid, text, index=0):
    info = m.text_artifact(pid, aid)
    return m.save_text_revision(
        pid,
        {
            "request_id": str(uuid.uuid4()),
            "parent_artifact_id": aid,
            "parent_sha256": info["sha256"],
            "head_id": info["head_id"],
            "head_sha256": info["head_sha256"],
            "locator": info["pages"][index]["locator"],
            "text": text,
        },
    )["artifact_id"]


def test_actual_revision_range_multi_fragment_and_legacy_html(project):
    m, pid, aid = project
    root = m.path(pid)
    old = m.get(pid)
    artifact = next(a for a in old["artifacts"] if a["artifact_id"] == aid)
    resource = root / artifact["pages"][0]["data"]
    before = resource.read_bytes()
    assert "reading_html" not in json.loads(before.decode().split("=", 1)[1].rstrip(";\n"))
    fallback = json.loads((root / "review/reading-fallbacks.js").read_text().split("=", 1)[1].rstrip(";"))
    assert fallback["来源第6页完整文字。"] == "<p>来源第6页完整文字。</p>"
    text = "# 标题\n\n普通 **加粗** 与 *斜体*，①②③ ⸺ 📖。\n保留换行和普通 Page 99 文字。\n\n- 第一项\n- 第二项\n\n3. 原序号三\n7. 原序号七"
    new = revise(m, pid, aid, text)
    job = m.export_documents(pid, {"source_artifact_id": new, "pages": "6-10", "formats": ["docx", "pdf"]})
    p, task = terminal(m, pid, job["task_id"])
    assert task["status"] == "SUCCEEDED", task.get("error")
    outputs = [a for a in p["artifacts"] if a.get("task_id") == job["task_id"] and a["role"] == "export"]
    word = next(a for a in outputs if a["format"] == "docx")
    doc = Document(root / word["path"])
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "普通 加粗 与 斜体，①②③ ⸺ 📖。\n保留换行和普通 Page 99 文字。" in body
    assert "来源第7页的第二片段必须保留。" in body
    assert "3. 原序号三" in body and "7. 原序号七" in body
    assert not doc.tables and any(r.bold and r.text == "加粗" for p in doc.paragraphs for r in p.runs)
    assert any(r.italic and r.text == "斜体" for p in doc.paragraphs for r in p.runs)
    normal = next(p for p in doc.paragraphs if p.text.startswith("普通"))
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert abs(normal.paragraph_format.line_spacing.pt - 10.5 * 1.85) < 0.05
    pdf = next(a for a in outputs if a["format"] == "pdf")
    with fitz.open(root / pdf["path"]) as rendered:
        assert len(rendered) == 5
        assert "第二片段" in rendered[1].get_text() and "第二片段" not in rendered[0].get_text()
    assert [g["source_page"] for g in word["source_output_mapping"]] == [6, 7, 8, 9, 10]
    assert word["export_source"]["artifact_id"] == new
    assert word["export_source"]["mode"] == "markdown-direct"
    assert digest(root / artifact["path"]) == artifact["sha256"] and resource.read_bytes() == before


@pytest.mark.parametrize(
    "mutation,match",
    [
        (lambda v: v.pop("source_page"), "映射"),
        (lambda v: v.update(source_pages=[6, 7]), "多个页"),
        (lambda v: v.pop("data"), "文字记录资源"),
    ],
)
def test_no_mapping_guess(project, mutation, match):
    m, pid, aid = project
    p = m.get(pid)
    a = copy.deepcopy(next(a for a in p["artifacts"] if a["artifact_id"] == aid))
    mutation(a["pages"][0])
    with pytest.raises(ValueError, match=match):
        export_options({"mode": "markdown-direct"}, p, a)


@pytest.mark.parametrize(
    "text,match",
    [
        ("", "没有正文"),
        ("![图片](private.png)", "不支持"),
        ("# 密集页\n\n" + "这是一段不能截断的正文。" * 1800, "来源第6页"),
    ],
)
def test_empty_unsupported_and_overflow_never_publish(project, text, match):
    m, pid, aid = project
    new = revise(m, pid, aid, text)
    job = m.export_documents(
        pid, {"mode": "markdown-direct", "source_artifact_id": new, "pages": "6", "formats": ["pdf"]}
    )
    p, t = terminal(m, pid, job["task_id"])
    assert t["status"] == "FAILED", t
    assert match in t["error"]["message"]
    assert not any(a.get("task_id") == job["task_id"] for a in p["artifacts"])


def test_direct_bypasses_question_parser_and_adapts_only_when_allowed(tmp_path):
    # Deliberately exceeds A4 at the actual CSS line-height, but fits at the
    # explicitly requested compact boundary; do not depend on Word font metrics.
    text = "# 阅读标题\n\n" + "\n\n".join("正常长段落。" * 16 for _ in range(20))
    template = direct_template()
    req = {
        "export_mode": "markdown-direct",
        "output": str(tmp_path),
        "source_artifact_id": "synthetic",
        "source_sha256": "fixed",
        "template_data": template,
        "pages": [
            {
                "source_pdf_physical_page": 23,
                "text": text,
                "source_locator": {"kind": "block", "id": "stable"},
            }
        ],
    }
    with patch(
        "nas_filetools.standalone.documents.structure_ocr", side_effect=AssertionError("must not infer")
    ):
        with pytest.raises(ValueError, match="来源第23页"):
            generate(req)
        compact = tmp_path / "compact"
        compact.mkdir()
        req["output"] = str(compact)
        req["direct_compact"] = {
            "enabled": True,
            "minimum_font_pt": 8,
            "minimum_line_multiple": 1.2,
            "minimum_paragraph_em": 0,
        }
        generate(req)
    report = json.loads((compact / "pagination-report.json").read_text())
    assert report["verified"] and len(report["groups"][0]["attempts"]) > 1


def test_failed_extraction_is_not_successful_blank_page(project):
    m, pid, aid = project
    with m.lock:
        p = m.get(pid)
        a = next(a for a in p["artifacts"] if a["artifact_id"] == aid)
        a["pages"][0]["error"] = True
        m.save(m.path(pid), p)
    job = m.export_documents(pid, {"source_artifact_id": aid, "pages": "6", "formats": ["pdf"]})
    p, task = terminal(m, pid, job["task_id"])
    assert task["status"] == "FAILED"
    assert "来源第6页提取失败" in task["error"]["message"]
    assert not any(a.get("task_id") == job["task_id"] for a in p["artifacts"])
